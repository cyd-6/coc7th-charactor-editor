from __future__ import annotations

import hashlib
import re
import warnings
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import OccupationDefinition, SkillChoiceGroup, SkillDefinition
from .currency import load_exchange_data


GROUP_LABELS = {
    "☆": "第一选择组",
    "⊙": "第二选择组",
    "☯": "社交技能",
    "※": "多选技能",
}


@dataclass(frozen=True, slots=True)
class ExcelSkillCell:
    template_slot: str
    name_cell: str
    specialization_cell: str
    base_cell: str
    extra_cell: str
    experience_cell: str
    occupation_cell: str
    interest_cell: str
    total_cell: str


@dataclass(frozen=True, slots=True)
class WeaponDefinition:
    catalog_id: str
    name: str
    group: str
    skill: str
    damage: str
    range: str
    attacks: str
    ammo: str
    malfunction: str
    era: str


@dataclass(frozen=True, slots=True)
class InventoryDefinition:
    catalog_id: str
    name: str
    group: str
    era: str
    price: str
    pack: str = ""


@dataclass(frozen=True, slots=True)
class TemplateCatalog:
    template_path: Path
    source_sha256: str
    sheet_names: tuple[str, ...]
    occupations: tuple[OccupationDefinition, ...]
    skills: tuple[SkillDefinition, ...]
    skill_cells: tuple[ExcelSkillCell, ...]
    weapons: tuple[WeaponDefinition, ...] = ()
    inventory: tuple[InventoryDefinition, ...] = ()
    experience_packages: tuple[dict, ...] = ()
    currency_quotes: tuple[dict, ...] = ()
    currency_metadata: dict | None = None
    template_revision_note: str | None = None

    @classmethod
    def load(cls, template_path: str | Path) -> "TemplateCatalog":
        path = Path(template_path).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"找不到调查员卡模板：{path}")

        source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Data Validation extension is not supported and will be removed",
            )
            # Stream each required region once. Random cell() calls in read-only
            # mode would repeatedly parse the XML, so readers use bounded rows.
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)

        try:
            exchange_data = load_exchange_data()
            version = workbook.defined_names.get("COC7_FX_DATA_SHA256")
            if version is None or version.attr_text.strip('"') != exchange_data["data_sha256"]:
                raise ValueError("模板与年度汇率数据版本不一致，请重新生成模板后重启服务。")
            points_version = workbook.defined_names.get("COC7_SKILL_POINTS_SCHEMA")
            if points_version is None or points_version.attr_text != '"separate-v1"':
                raise ValueError("模板尚未区分经历包点与成长点数，请更新 XLSX 模板后重启服务。")
            occupations = _load_occupations(workbook)
            skills, skill_cells = _load_skills(workbook)
            revision_note = next(workbook['更新说明'].iter_rows(
                min_row=3, max_row=3, min_col=16, max_col=16, values_only=True))[0]
            if revision_note is not None and (not isinstance(revision_note,str) or revision_note.startswith('=')):
                raise ValueError('模板版本说明必须为文本或空白：更新说明!P3')
            return cls(
                template_path=path,
                source_sha256=source_hash,
                sheet_names=tuple(workbook.sheetnames),
                occupations=tuple(occupations),
                skills=tuple(skills),
                skill_cells=tuple(skill_cells),
                template_revision_note=revision_note,
                weapons=tuple(_load_weapons(workbook)),
                inventory=tuple(_load_inventory(workbook)),
                currency_quotes=tuple(exchange_data["quotes"]),
                currency_metadata={k: v for k, v in exchange_data.items() if k != "quotes"},
                experience_packages=tuple({
                    "name": str(values[0]),
                    "san_description": str(values[1]),
                    "skill_points": int(values[2]),
                    "notes": str(values[3]).strip(),
                    "minimum_age": {21: 25, 22: 20, 23: 30}.get(row, 0),
                } for row, values in enumerate(workbook["附表"].iter_rows(
                    min_row=20, max_row=24, min_col=2, max_col=5, values_only=True), 20)),
            )
        finally:
            workbook.close()

    def occupation_by_id(self, occupation_id: int) -> OccupationDefinition | None:
        return next((item for item in self.occupations if item.occupation_id == occupation_id), None)

    def occupation_by_name(self, name: str) -> OccupationDefinition | None:
        target = name.strip().casefold()
        return next((item for item in self.occupations if item.name.strip().casefold() == target), None)

    def skill_cell(self, template_slot: str) -> ExcelSkillCell | None:
        return next((item for item in self.skill_cells if item.template_slot == template_slot), None)


def _literal_text(value: object) -> str:
    if value is None or isinstance(value, str) and value.startswith("="):
        return ""
    return str(value).strip()


def _load_weapons(workbook) -> list[WeaponDefinition]:
    sheet = workbook["武器列表 战斗"]
    result = []
    group = ""
    # 人物卡!G54:L58 validates against B2:B105. B contains era-dependent
    # display formulas; U holds their literal names, independent of Excel caches.
    for row, values in enumerate(sheet.iter_rows(min_row=2, max_row=105, max_col=21, values_only=True), 2):
        read = lambda col: _literal_text(values[col - 1])
        group = read(13).replace("\n", "") or group
        if read(21):
            result.append(WeaponDefinition(
                catalog_id=f"weapon:{row}", name=read(21), group=group,
                skill=read(3), damage=read(4), range=read(5),
                attacks=read(7), ammo=read(8), malfunction=read(9), era=read(10),
            ))
    return result


def _load_inventory(workbook) -> list[InventoryDefinition]:
    sheet = workbook["资产及物价参考"]
    # Explicit goods sections exclude fares, rent, tuition and table headings.
    # (era, category, first row, last row, name column, price column, pack column)
    sections = (
        ("现代", "男装", 45, 62, 2, 5, 0),
        ("现代", "女装", 66, 82, 2, 5, 0),
        ("现代", "帐篷", 71, 72, 7, 11, 0),
        ("现代", "医疗用品", 78, 82, 7, 11, 0),
        ("现代", "户外与旅行装备", 45, 60, 13, 16, 0),
        ("现代", "行李", 64, 66, 13, 16, 0),
        ("现代", "通讯设备", 46, 47, 18, 20, 0),
        ("现代", "计算机", 51, 55, 18, 20, 0),
        ("现代", "工具", 60, 63, 18, 20, 0),
        ("现代", "电子器材", 67, 83, 18, 20, 0),
        ("现代", "子弹", 45, 59, 22, 27, 26),
        ("现代", "火器配件与近战装备", 63, 74, 22, 26, 0),
        ("1920s", "男装", 93, 125, 2, 4, 0),
        ("1920s", "女装", 129, 153, 2, 4, 0),
        ("1920s", "户外服饰", 157, 168, 2, 4, 0),
        ("1920s", "行李", 172, 177, 2, 4, 0),
        ("1920s", "通讯设备", 185, 189, 2, 4, 0),
        ("1920s", "户外与旅行装备", 93, 119, 7, 10, 0),
        ("1920s", "利刃", 121, 123, 7, 10, 0),
        ("1920s", "陷阱", 125, 127, 7, 10, 0),
        ("1920s", "帐篷", 93, 99, 13, 16, 0),
        ("1920s", "液体容器", 101, 105, 13, 16, 0),
        ("1920s", "医疗用品", 109, 127, 13, 16, 0),
        ("1920s", "化妆品与卫生用品", 131, 142, 13, 16, 0),
        ("1920s", "汽车配件", 179, 188, 13, 16, 0),
        ("1920s", "留声机与照相机", 198, 206, 13, 16, 0),
        ("1920s", "乐器", 208, 216, 13, 16, 0),
        ("1920s", "酒水", 93, 98, 19, 21, 0),
        ("1920s", "食材", 102, 121, 19, 21, 0),
        ("1920s", "工具", 132, 148, 19, 21, 0),
        ("1920s", "调查工具", 152, 185, 19, 21, 0),
        ("1920s", "运动与游戏", 189, 212, 19, 21, 0),
        ("1920s", "子弹", 93, 109, 24, 27, 26),
        ("1920s", "枪械配件", 113, 115, 24, 26, 0),
        ("1920s", "近战装备", 119, 127, 24, 26, 0),
    )
    result = []
    first_row = min(section[2] for section in sections)
    rows = tuple(sheet.iter_rows(min_row=first_row, max_row=max(s[3] for s in sections),
                                 max_col=27, values_only=True))
    for era, group, first, last, name_col, price_col, pack_col in sections:
        for row in range(first, last + 1):
            values = rows[row - first_row]
            name = _literal_text(values[name_col - 1])
            price = _literal_text(values[price_col - 1])
            if not name or not price:
                continue
            result.append(InventoryDefinition(
                catalog_id=f"inventory:{get_column_letter(name_col)}{row}", name=name,
                group=group, era=era, price=price,
                pack=_literal_text(values[pack_col - 1]) if pack_col else "",
            ))
    return result


def normalize_skill_name(value: object) -> str:
    text = str(value or "").replace(" Ω", "").replace("Ω", "").strip()
    return re.sub(r"\s+", " ", text)


def occupation_skill_token(slot_name: str, specialization: str = "") -> str:
    slot = normalize_skill_name(slot_name)
    spec = specialization.strip()
    return f"{slot}（{spec}）" if spec else slot


def split_occupation_skill_token(token: str) -> tuple[str, str]:
    match = re.fullmatch(r"(.+?)（(.+)）", token.strip())
    if match:
        return normalize_skill_name(match.group(1)), match.group(2).strip()
    return normalize_skill_name(token), ""


def _parse_credit_range(value: object) -> tuple[int, int]:
    text = str(value or "").strip()
    match = re.search(r"(\d+)\s*[-—~～至]\s*(\d+)", text)
    if not match:
        return 0, 99
    low, high = int(match.group(1)), int(match.group(2))
    return (low, high) if low <= high else (high, low)


def _load_occupations(workbook) -> list[OccupationDefinition]:
    sheet = workbook["职业列表"]
    matrix = workbook["本职技能"]
    matrix_rows = tuple(matrix.iter_rows(max_row=min(matrix.max_row, 72), values_only=True))
    id_to_column: dict[int, int] = {}
    for column, value in enumerate(matrix_rows[0][1:], 1):
        if isinstance(value, (int, float)):
            id_to_column[int(value)] = column

    occupations: list[OccupationDefinition] = []
    for values in sheet.iter_rows(min_row=3, max_row=min(sheet.max_row, 232), max_col=13, values_only=True):
        occupation_id = values[0]
        name = values[1]
        if not isinstance(occupation_id, (int, float)) or int(occupation_id) <= 1:
            continue
        if not isinstance(name, str) or not name.strip() or name.startswith("="):
            continue

        occupation_id = int(occupation_id)
        credit_min, credit_max = _parse_credit_range(values[3])
        point_formula = str(values[5] or "=EDU*4").lstrip("=")
        matrix_column = id_to_column.get(occupation_id)
        fixed_skills: list[str] = []
        group_members: dict[str, list[str]] = {marker: [] for marker in GROUP_LABELS}
        group_counts: dict[str, int] = {marker: 0 for marker in GROUP_LABELS}
        free_choices = 0

        if matrix_column:
            for marker_row, marker in zip(range(3, 7), GROUP_LABELS, strict=True):
                value = matrix_rows[marker_row - 1][matrix_column]
                if isinstance(value, (int, float)):
                    group_counts[marker] = max(0, int(value))
            free_value = matrix_rows[6][matrix_column]
            if isinstance(free_value, (int, float)):
                free_choices = max(0, int(free_value))

            for matrix_values in matrix_rows[7:]:
                slot_name = normalize_skill_name(matrix_values[0])
                marker_value = matrix_values[matrix_column]
                if not slot_name or marker_value in (None, "", 0):
                    continue
                marker_text = str(marker_value).strip()
                if marker_text in GROUP_LABELS:
                    group_members[marker_text].append(slot_name)
                elif marker_text == "★":
                    fixed_skills.append(slot_name)
                else:
                    fixed_skills.append(occupation_skill_token(slot_name, marker_text))

        choice_groups = tuple(
            SkillChoiceGroup(
                marker=marker,
                label=GROUP_LABELS[marker],
                required_count=group_counts[marker],
                candidates=tuple(dict.fromkeys(group_members[marker])),
            )
            for marker in GROUP_LABELS
            if group_counts[marker] or group_members[marker]
        )
        occupations.append(
            OccupationDefinition(
                occupation_id=occupation_id,
                name=name.strip(),
                credit_min=credit_min,
                credit_max=credit_max,
                point_formula=point_formula,
                summary=str(values[6] or "").strip(),
                contacts=str(values[10] or "").strip(),
                description=str(values[12] or "").strip(),
                fixed_skills=tuple(dict.fromkeys(fixed_skills)),
                choice_groups=choice_groups,
                free_choices=free_choices,
            )
        )
    return occupations


def _base_value_and_formula(name: str, raw_value: object, specialization: str) -> tuple[int, str | None]:
    if isinstance(raw_value, (int, float)):
        return int(raw_value), None
    formula = str(raw_value or "")
    upper = formula.upper()
    if "DEX" in upper:
        return 0, "DEX/2"
    if "AG5" in upper or name == "母语":
        return 0, "EDU"
    if name.startswith("技艺"):
        return 5, None
    if name.startswith("格斗"):
        return (25 if specialization == "斗殴" else 1), None
    if name.startswith("射击"):
        return (20 if specialization == "手枪" else 1), None
    if name.startswith("科学"):
        return (10 if specialization == "数学" else 1), None
    if name.startswith("生存"):
        return 10, None
    return 1, None


def _load_skills(workbook) -> tuple[list[SkillDefinition], list[ExcelSkillCell]]:
    sheet = workbook["人物卡"]
    skills: list[SkillDefinition] = []
    cells: list[ExcelSkillCell] = []
    sides = (
        ("left", 6, 8, 10, 12, 14, 16, 18),
        ("right", 28, 30, 32, 34, 36, 38, 40),
    )
    for row, values in enumerate(sheet.iter_rows(min_row=16, max_row=49, max_col=40, values_only=True), 16):
        for side, name_col, spec_col, base_col, extra_col, occ_col, int_col, total_col in sides:
            raw_name = values[name_col - 1]
            if not isinstance(raw_name, str) or not raw_name.strip() or raw_name.startswith("="):
                continue
            name = normalize_skill_name(raw_name)
            specialization = str(values[spec_col - 1] or "").strip()
            base_value, base_formula = _base_value_and_formula(
                name,
                values[base_col - 1],
                specialization,
            )
            address = lambda column: f"{get_column_letter(column)}{row}"
            template_slot = address(name_col)
            specializable = bool(specialization) or name.endswith(("：", ":")) or bool(re.search(r"[①②③]$", name))
            skills.append(
                SkillDefinition(
                    key=template_slot,
                    name=name,
                    base_value=base_value,
                    template_slot=template_slot,
                    base_formula=base_formula,
                    specializable=specializable,
                    default_specialization=specialization,
                )
            )
            cells.append(
                ExcelSkillCell(
                    template_slot=template_slot,
                    name_cell=address(name_col),
                    specialization_cell=address(spec_col),
                    base_cell=address(base_col),
                    extra_cell=address(extra_col),
                    experience_cell=address(extra_col + 1),
                    occupation_cell=address(occ_col),
                    interest_cell=address(int_col),
                    total_cell=address(total_col),
                )
            )
    return skills, cells
