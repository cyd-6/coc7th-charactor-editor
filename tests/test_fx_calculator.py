"""Top-of-sheet cross-currency calculator and independent asset selector."""
import json
import os
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import get_catalog
from coc7_card.exporters.xlsx_template import TemplateWorkbook
from coc7_card.importers.excel import import_investigator
from coc7_card.template_config import TEMPLATE_PATH

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = TEMPLATE_PATH
QUOTES = json.loads((ROOT/'assets/rates/annual.json').read_text(encoding='utf-8'))['quotes']


def quote(year, code):
    return next(q for q in QUOTES if q['year']==year and q['code']==code)


def test_calculator_layout_and_independent_asset_authority():
    book = load_workbook(TEMPLATE)
    try:
        fx, card = book['货币汇率'], book['人物卡']
        assert fx['A1'].value == '货币换算'
        assert fx['C7'].data_type == 'f'
        assert fx['K2'].value.replace("'",'').replace('$','') == '=人物卡!S62'
        assert card['S62'].data_type != 'f'
        assert 'X36:AE36' in fx.merged_cells
        assert fx['X39'].value == 1920
        assert fx.sheet_view.topLeftCell == 'A3'
        assert [fx[f'A{row}'].value for row in range(3,8)] == ['当前年份','原始金额','原币种','目标币种','目标金额']
        assert all(fx.row_dimensions[row].hidden for row in [1,2,8,9])
        assert all(not fx.row_dimensions[row].hidden for row in range(3,8))
        for row in range(35, 1157):
            assert not fx.row_dimensions[row].hidden
            assert not fx.row_dimensions[row].collapsed
            assert fx.row_dimensions[row].outlineLevel == 0
        assert {fx.cell(row, 1).value for row in range(36, 1157)} == set(range(1920, 2027))
        assert all(not c.hidden and not c.collapsed and not c.outlineLevel
                   for c in fx.column_dimensions.values() if c.min <= 20)
        assert all(c.hidden for c in fx.column_dimensions.values() if c.min >= 21)
        assert all(fx[f'{col}{row}'].number_format == ';;;'
                   for col in ('J', 'K') for row in range(3, 8))
        assert book.defined_names['COC7_FX_ASSET_OPTIONS'].attr_text == 'INDIRECT("FX_Y"&\'货币汇率\'!$K$1)'
        assert fx.print_area == "'货币汇率'!$A$3:$F$7"
    finally:
        book.close()


def test_import_prefers_main_currency_input_over_old_formula_cache(tmp_path):
    book = TemplateWorkbook(TEMPLATE)
    try:
        book.Worksheets('货币汇率').Range('K1').Value2 = 1930
        book.Worksheets('人物卡').Range('S62').Value2 = quote(1930,'GBP')['name']
        # K2's saved result is intentionally not recalculated.
        output = tmp_path/'edited-currency.xlsx'
        book.save(output)
    finally:
        book.close()
    result = import_investigator(output.read_bytes(), get_catalog())
    assert result['draft']['assets']['currency_code'] == 'GBP'
    assert result['draft']['assets']['exchange_year'] == 1930


def test_import_legacy_currency_input_remains_supported(tmp_path):
    book = TemplateWorkbook(TEMPLATE)
    try:
        fx, card = book.Worksheets('货币汇率'), book.Worksheets('人物卡')
        # CY26.2 kept its editable selector in K2 instead of the main card.
        fx.Range('K1').Value2 = 1930
        fx.Range('K2').Value2 = quote(1930, 'GBP')['name']
        fx.Range('K10').Value2 = 123.45
        fx.Range('K11').Value2 = 0
        card.Range('S62').Formula = '=BK37'
        book.Worksheets('更新说明').Range('P3').Value2 = '版本号CY26.2'
        output = tmp_path / 'legacy-currency.xlsx'
        book.save(output)
    finally:
        book.close()
    assets = import_investigator(output.read_bytes(), get_catalog())['draft']['assets']
    assert assets['currency_code'] == 'GBP'
    assert assets['exchange_year'] == 1930
    assert float(assets['usd_cash']) == 123.45
    assert assets['usd_assets'] != '' and float(assets['usd_assets']) == 0


@pytest.mark.skipif(sys.platform != 'win32' or os.environ.get('COC7_TEST_EXCEL') != '1',
                    reason='Opt-in native Excel validation')
def test_native_cross_conversion_missing_inputs_and_independence(tmp_path):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    excel = book = fx = card = None
    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3
        book = excel.Workbooks.Open(str(TEMPLATE), UpdateLinks=0, ReadOnly=True, CorruptLoad=0)
        fx, card = book.Worksheets('货币汇率'), book.Worksheets('人物卡')
        for cell in ('C5','C6'):
            assert fx.Range(cell).Validation.Type == 3
        assert card.Range('S62').Validation.Type == 3
        for year, source, target, amount in [
            (1930,'JPY','GBP',100), (1930,'JPY','JPY',123.45),
            (2026,'USD','GBP',0), (1930,'GBP','USD',-5),
            (1954,'CNY_OLD','USD',100), (1955,'CNY','USD',100),
            (1959,'FRF_OLD','USD',100), (1960,'FRF','USD',100),
            (1997,'RUR','USD',100), (1998,'RUB','USD',100),
            (1999,'EUR','USD',100), (2026,'GBP','JPY',100),
            (1923,'DE_TRANSITION_1923','USD',100), (1923,'USD','DE_TRANSITION_1923',100),
        ]:
            a, b = quote(year,source), quote(year,target)
            for cell,value in [('C3',year),('C4',amount),('C5',a['name']),('C6',b['name'])]:
                fx.Range(cell).Value2 = value
            if a['available'] and b['available']:
                expected = round(amount/float(a['rate'])*float(b['rate']),2)
                assert fx.Range('C7').Value2 == pytest.approx(expected,abs=.0001)
                assert fx.Range('C9').Value2 == '可用'
            else:
                assert fx.Range('C7').Value2 == fx.Range('C9').Value2
                assert '缺少报价' in fx.Range('C9').Value2
        fx.Range('C3').Value2 = 1930
        fx.Range('C5').Value2 = quote(1930,'JPY')['name']
        fx.Range('C6').Value2 = quote(1930,'GBP')['name']
        for bad in ['',None,'abc']:
            fx.Range('C4').Value2 = bad
            assert fx.Range('C7').Value2 == fx.Range('C9').Value2
            assert '数值金额' in fx.Range('C9').Value2
        fx.Range('C4').Value2 = 100
        for bad in [1919,2027,1930.5,'bad',None]:
            fx.Range('C3').Value2 = bad
            assert fx.Range('C7').Value2 == fx.Range('C9').Value2
            assert '年份' in fx.Range('C9').Value2
        fx.Range('C3').Value2 = 1998
        fx.Range('C5').Value2 = quote(1999,'EUR')['name']
        assert fx.Range('C7').Value2 == fx.Range('C9').Value2
        assert '原币种' in fx.Range('C9').Value2
        fx.Range('C3').Value2 = 1930
        fx.Range('C5').Value2 = quote(1930,'JPY')['name']
        fixed = fx.Range('C7').Value2
        fx.Range('K10').Value2 = 123.45
        for year,code in [(1930,'GBP'),(2026,'CNY'),(1923,'DE_TRANSITION_1923')]:
            q = quote(year,code)
            fx.Range('K1').Value2 = year
            card.Range('S62').Value2 = q['name']
            assert fx.Range('C7').Value2 == fixed
            assert fx.Range('K10').Value2 == 123.45
            if q['available']:
                assert card.Range('O62').Value2 == pytest.approx(round(123.45*float(q['rate']),2))
                options = fx.Evaluate('COC7_FX_ASSET_OPTIONS').Value2
                assert q['name'] in [row[0] for row in options]
            else:
                assert card.Range('O62').Value2 in ('',None)
        fx.Range('K1').Value2 = 1930
        card.Range('S62').Value2 = quote(1930,'GBP')['name']
        saved = tmp_path/'native-saved.xlsx'
        book.SaveCopyAs(str(saved))
        book.Close(SaveChanges=False)
        book = excel.Workbooks.Open(str(saved),UpdateLinks=0,ReadOnly=True,CorruptLoad=0)
        assert book.Worksheets('人物卡').Range('S62').Validation.Type == 3
        assert book.Worksheets('货币汇率').Range('C7').Value2 == fixed
        imported = import_investigator(saved.read_bytes(),get_catalog())
        assert imported['draft']['assets']['currency_code'] == 'GBP'
    finally:
        if book is not None: book.Close(SaveChanges=False)
        if excel is not None: excel.Quit()
        card = fx = book = excel = None
        pythoncom.CoUninitialize()
