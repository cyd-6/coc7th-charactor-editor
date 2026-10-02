"""Preserve XLSX layout and formula caches without duplicate workbook objects."""

from __future__ import annotations

import gc
import io
import warnings
import weakref
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from coc7_card.exporters.excel import ExcelExportError
from coc7_card.exporters.excel_linux import LinuxExcelExporter
from coc7_card.exporters.xlsx_template import DRAWING, MAIN, PACKAGE, REL, TemplateWorkbook, tag
from coc7_card.template_config import TEMPLATE_PATH


def test_xml_layout_and_shared_formulas_match_original_template():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        reference = load_workbook(TEMPLATE_PATH, data_only=False)
    workbook = TemplateWorkbook(TEMPLATE_PATH)
    try:
        assert tuple(workbook.paths) == tuple(reference.sheetnames)
        assert len(workbook.sheets) == 14
        for name, sheet in workbook.sheets.items():
            original = reference[name]
            assert bool(sheet.PrintArea) == bool(original.print_area)
            for merged in original.merged_cells.ranges:
                assert sheet.Range(merged.start_cell.coordinate).MergeArea.address == str(merged)
            for address, cell in sheet.cells.items():
                formula = cell.find(tag("f"))
                if formula is not None:
                    assert "=" + formula.text == original[address].value, (name, address)
            # Include the portrait box and an unconfigured column/row so both
            # explicit dimensions and sheet defaults are compared.
            for column in (*range(38, 46), 200):
                widths = [dimension.width for dimension in original.column_dimensions.values()
                          if dimension.min <= column <= dimension.max]
                expected = widths[0] if widths else original.sheet_format.defaultColWidth or 8.43
                assert sheet.column_pixels(column) == expected * 7 + 5
            for row in (*range(3, 10), 30000):
                expected = original.row_dimensions[row].height or original.sheet_format.defaultRowHeight or 15
                assert sheet.row_pixels(row) == expected * 4 / 3
    finally:
        workbook.close()
        reference.close()


def test_xml_portrait_placement_preserves_original_box(tmp_path):
    workbook = TemplateWorkbook(TEMPLATE_PATH)
    try:
        sheet = workbook.Worksheets("人物卡")
        output = io.BytesIO()
        with Image.new("RGBA", (80, 120), (10, 20, 30, 100)) as image:
            image.save(output, format="PNG")
        workbook.add_portrait(sheet, output.getvalue())
        anchor = next(anchor for root in workbook.trees.values()
                      for anchor in root.iter(f"{{{DRAWING}}}oneCellAnchor")
                      if anchor.find(f".//{{{DRAWING}}}cNvPr[@name='调查员头像']") is not None)
        origin = anchor.find(f"{{{DRAWING}}}from")
        assert origin.find(f"{{{DRAWING}}}col").text == "37"
        assert origin.find(f"{{{DRAWING}}}row").text == "2"
        box_width = sum(sheet.column_pixels(column) for column in range(38, 46))
        box_height = sum(sheet.row_pixels(row) for row in range(3, 10))
        scale = min(box_width / 80, box_height / 120)
        extent = anchor.find(f"{{{DRAWING}}}ext")
        assert int(extent.get("cx")) == round(80 * scale * 9525)
        assert int(extent.get("cy")) == round(120 * scale * 9525)
        assert int(origin.find(f"{{{DRAWING}}}colOff").text) == round((box_width - 80 * scale) / 2 * 9525)
        assert int(origin.find(f"{{{DRAWING}}}rowOff").text) == round((box_height - 120 * scale) / 2 * 9525)
        path = tmp_path / "portrait.xlsx"
        workbook.save(path)
        with ZipFile(path) as archive:
            with Image.open(io.BytesIO(archive.read("xl/media/coc7-investigator-portrait.png"))) as image:
                assert image.size == (80, 120)
                assert image.getpixel((0, 0)) == (10, 20, 30, 100)
    finally:
        workbook.close()


def test_xml_layout_reads_updates_and_default_dimensions():
    workbook = TemplateWorkbook(TEMPLATE_PATH)
    try:
        sheet = workbook.Worksheets("人物卡")
        sheet.Columns("AL").ColumnWidth = 12.5
        assert sheet.column_pixels(38) == 12.5 * 7 + 5
        sheet.rows[3].set("ht", "20")
        assert sheet.row_pixels(3) == 20 * 4 / 3
        sheet.PrintArea = "$A$1:$D$8"
        assert sheet.PrintArea == "'人物卡'!$A$1:$D$8"
        merges = sheet.root.find(tag("mergeCells"))
        ET.SubElement(merges, tag("mergeCell"), ref="B200:D201")
        assert sheet.Range("C201").MergeArea.address == "B200:D201"
        assert sheet.Range("A200:D201").MergeArea.address == "A200:D201"
        sheet.root.remove(sheet.root.find(tag("sheetFormatPr")))
        assert sheet.row_pixels(30000) == 20
        sheet.root.remove(sheet.root.find(tag("cols")))
        assert sheet.column_pixels(200) == 8.43 * 7 + 5
    finally:
        workbook.close()


def test_closed_workbook_releases_sheets_without_cyclic_gc():
    was_enabled = gc.isenabled()
    gc.disable()
    try:
        workbook = TemplateWorkbook(TEMPLATE_PATH)
        book_ref = weakref.ref(workbook)
        sheet_refs = [weakref.ref(sheet) for sheet in workbook.sheets.values()]
        workbook.close()
        workbook.close()  # Cleanup stays safe in nested finally blocks.
        assert not workbook.parts and not workbook.trees
        assert not workbook.shared_strings
        del workbook
        assert book_ref() is None
        assert all(reference() is None for reference in sheet_refs)
    finally:
        if was_enabled:
            gc.enable()


@pytest.fixture
def cache_package(tmp_path):
    name, part = "测试", "xl/worksheets/sheet1.xml"
    root = ET.Element(tag("workbook"), nsmap={None: MAIN, "r": REL})
    sheets = ET.SubElement(root, tag("sheets"))
    ET.SubElement(sheets, tag("sheet"), name=name, sheetId="1", attrib={f"{{{REL}}}id": "rId1"})
    relations = ET.Element(f"{{{PACKAGE}}}Relationships")
    ET.SubElement(relations, f"{{{PACKAGE}}}Relationship", Id="rId1", Target="worksheets/sheet1.xml")
    calculated = ET.Element(tag("worksheet"), nsmap={None: MAIN})
    row = ET.SubElement(ET.SubElement(calculated, tag("sheetData")), tag("row"), r="1")
    strings = ET.Element(tag("sst"))
    ET.SubElement(ET.SubElement(strings, tag("si")), tag("t")).text = "共享文字"
    cells = {}
    for column, (kind, value) in enumerate([(None, "42"), ("s", "0"), ("inlineStr", "内联文字"), ("str", None), ("b", "1")], 1):
        address = f"{get_column_letter(column)}1"
        original = ET.Element(tag("c"), r=address, s="7", t="str")
        ET.SubElement(original, tag("f")).text = 'IF(TRUE,42,0)'
        ET.SubElement(original, tag("v")).text = "stale"
        cells[address] = original
        cached = ET.SubElement(row, tag("c"), r=address)
        if kind:
            cached.set("t", kind)
        if kind == "inlineStr":
            ET.SubElement(ET.SubElement(cached, tag("is")), tag("t")).text = value
        elif value is not None:
            ET.SubElement(cached, tag("v")).text = value
    untouched = ET.Element(tag("c"), r="F1", t="inlineStr")
    ET.SubElement(ET.SubElement(untouched, tag("is")), tag("t")).text = "保留原值"
    cells["F1"] = untouched
    workbook = SimpleNamespace(paths={name: part}, sheets={name: SimpleNamespace(cells=cells)})
    path = tmp_path / "calculated.xlsx"

    def save():
        with ZipFile(path, "w") as archive:
            archive.writestr("xl/workbook.xml", ET.tostring(root))
            archive.writestr("xl/_rels/workbook.xml.rels", ET.tostring(relations))
            archive.writestr("xl/sharedStrings.xml", ET.tostring(strings))
            archive.writestr(part, ET.tostring(calculated))
            archive.writestr("xl/media/unrelated.png", b"This part must not be read")
        return path

    return SimpleNamespace(workbook=workbook, cells=cells, row=row, root=root, save=save)


def test_calculated_cache_merge_skips_unrelated_parts_and_preserves_formulas(cache_package, monkeypatch):
    fixture = cache_package
    path = fixture.save()
    opened = []
    original_open = ZipFile.open

    def track_open(self, name, *args, **kwargs):
        opened.append(name)
        return original_open(self, name, *args, **kwargs)

    monkeypatch.setattr(ZipFile, "open", track_open)
    before = ET.tostring(fixture.cells["F1"])
    LinuxExcelExporter._merge_calculated_values(fixture.workbook, path)
    assert "xl/media/unrelated.png" not in opened
    assert ET.tostring(fixture.cells["F1"]) == before
    for address, kind, value in [("A1", None, "42"), ("B1", "str", "共享文字"),
                                 ("C1", "str", "内联文字"), ("D1", "str", None), ("E1", "b", "1")]:
        cell = fixture.cells[address]
        assert cell.get("t") == kind
        assert cell.find(tag("v")).text == value
        assert cell.get("s") == "7"
        assert cell.find(tag("f")).text == 'IF(TRUE,42,0)'


@pytest.mark.parametrize("failure,message", [("missing_cell", "缺少重算结果"),
                                            ("missing_value", "缺少缓存值"),
                                            ("error", "公式错误"),
                                            ("sheet_name", "数量或顺序")])
def test_calculated_cache_merge_keeps_validation(cache_package, failure, message):
    fixture = cache_package
    cell = fixture.row[0]
    if failure == "missing_cell":
        fixture.row.remove(cell)
    elif failure == "missing_value":
        cell.remove(cell.find(tag("v")))
    elif failure == "error":
        cell.set("t", "e")
        cell.find(tag("v")).text = "#DIV/0!"
    else:
        fixture.root.find(tag("sheets"))[0].set("name", "错误工作表")
    with pytest.raises(ExcelExportError, match=message):
        LinuxExcelExporter._merge_calculated_values(fixture.workbook, fixture.save())
