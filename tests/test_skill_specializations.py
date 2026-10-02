"""Grouped skill labels keep independent values across draft and export paths."""

import io
import shutil
import warnings

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import FONT_PATH, get_catalog
from coc7_card.exporters.excel import ExcelExporter
from coc7_card.exporters.excel_linux import LinuxExcelExporter
from coc7_card.exporters.pdf import PdfExporter
from coc7_card.exporters.xlsx_template import TemplateWorkbook
from coc7_card.importers.excel import import_investigator
from coc7_card.rules import RuleEngine
from coc7_card.web import build_draft, catalog_payload


@pytest.fixture(scope="module")
def catalog():
    return get_catalog()


@pytest.fixture
def payload(catalog):
    occupation = next(item for item in catalog.occupations if item.credit_min == 0)
    return {
        "identity": {"name": "分项验证调查员", "age": 30},
        "occupation_id": occupation.occupation_id,
        "skills": [
            {"template_slot": "F38", "specialization": "手枪", "interest_points": 11},
            {"template_slot": "F39", "specialization": "步枪/霰弹枪", "interest_points": 12},
            {"template_slot": "F40", "specialization": "冲锋枪", "interest_points": 13},
            {"template_slot": "F41", "specialization": "弓术", "interest_points": 14},
            {"template_slot": "AB31", "specialization": "数学", "interest_points": 15},
        ],
    }


def test_catalog_groups_presets_and_options_keep_template_identity(catalog):
    definitions = {item.template_slot: item for item in catalog.skills}
    skills = {item["template_slot"]: item for item in catalog_payload(catalog)["skills"]}
    assert set(skills) == set(definitions)
    expected = {
        "F20": ("技艺", "摄影", 5), "F21": ("技艺", "美术", 5), "F22": ("技艺", "写作", 5),
        "F34": ("格斗", "斗殴", 25), "F35": ("格斗", "斧", 15),
        "F36": ("格斗", "剑", 20), "F37": ("格斗", "矛", 20),
        "F38": ("射击", "手枪", 20), "F39": ("射击", "步枪/霰弹枪", 25),
        "F40": ("射击", "冲锋枪", 15), "F41": ("射击", "弓术", 15),
        "AB31": ("科学", "物理学", 1), "AB32": ("科学", "化学", 1),
        "AB33": ("科学", "生物学", 1),
    }
    for slot, (group, specialization, base) in expected.items():
        skill = skills[slot]
        assert skill["group"] == group
        assert skill["preset_specialization"] == specialization
        assert next(item["base_value"] for item in skill["specialization_options"]
                    if item["name"] == specialization) == base
        assert skill["name"] == definitions[slot].name
        assert skill["default_specialization"] == definitions[slot].default_specialization
    assert {"name": "数学", "base_value": 10, "aliases": []} in skills["AB31"]["specialization_options"]
    assert skills["F46"]["group"] == "外语"
    assert skills["F46"]["specialization_options"] == []
    assert skills["F49"]["group"] == skills["F30"]["group"] == ""


def test_subskills_keep_independent_bases_and_allocations(catalog, payload):
    draft = build_draft(payload, catalog)
    skills = {item.template_slot: item for item in draft.skills}
    assert [(skills[slot].base_value, skills[slot].final_value) for slot in
            ("F38", "F39", "F40", "F41", "AB31")] == [
        (20, 31), (25, 37), (15, 28), (15, 29), (10, 25),
    ]
    assert len(skills) == len(catalog.skills)
    assert skills["F39"].name == "射击①"
    assert skills["F40"].display_name == "射击②（冲锋枪）"
    assert skills["F40"].interest_points == 13
    assert skills["F40"].occupation_points == skills["F40"].extra_final == skills["F40"].experience_points == 0


def test_old_empty_draft_does_not_receive_presets(catalog):
    draft = build_draft({}, catalog)
    for definition, skill in zip(catalog.skills, draft.skills, strict=True):
        assert skill.name == definition.name
        assert skill.specialization == definition.default_specialization
        assert skill.base_value == RuleEngine.skill_base_value(
            definition.base_formula, draft.attributes, definition.base_value,
        )


@pytest.mark.parametrize("occupation_name,slot,specialization,base", [
    ("绅士、淑女", "F39", "步/霰", 25),
    ("工人-伐木工", "F35", "链锯", 10),
    ("士兵", "F39", "来复", 25),
])
def test_occupation_short_names_use_known_bases(catalog, occupation_name, slot, specialization, base):
    occupation = next(item for item in catalog.occupations if item.name == occupation_name)
    skill = next(item for item in build_draft({"occupation_id": occupation.occupation_id}, catalog).skills
                 if item.template_slot == slot)
    assert (skill.specialization, skill.base_value, skill.selected_occupation) == (specialization, base, True)


@pytest.mark.parametrize("override", [0, 7, 35])
def test_imported_base_override_wins_over_named_specialization(catalog, payload, override):
    payload["skills"][2]["base_override"] = override
    skill = next(item for item in build_draft(payload, catalog).skills if item.template_slot == "F40")
    assert skill.base_value == override
    assert skill.final_value == override + 13


@pytest.mark.parametrize("name,specialization", [("射击②", "自制投射器"), ("档案研究", "冲锋枪")])
def test_custom_skill_names_and_unknown_subskills_keep_fallback(catalog, name, specialization):
    skill = next(item for item in build_draft({"skills": [{
        "template_slot": "F40", "name": name, "specialization": specialization, "interest_points": 9,
    }]}, catalog).skills if item.template_slot == "F40")
    assert (skill.name, skill.specialization, skill.base_value, skill.final_value) == (name, specialization, 1, 10)


def test_named_subskills_reach_two_page_pdf(catalog, payload):
    draft = build_draft(payload, catalog)
    result = PdfExporter(FONT_PATH).export(draft, include_preview=False)
    reader = PdfReader(io.BytesIO(result.data))
    assert len(reader.pages) == result.page_count == 2
    text = "".join("".join(page.extract_text() for page in reader.pages).split())
    for slot in ("F38", "F39", "F40", "F41", "AB31"):
        skill = next(item for item in draft.skills if item.template_slot == slot)
        assert skill.display_name in text
        assert f"{skill.final_value}{skill.hard_value}{skill.extreme_value}{skill.display_name}" in text


@pytest.mark.parametrize("saved_base", [0, 1, 15])
def test_import_keeps_saved_base_even_when_it_matches_old_empty_slot(catalog, payload, tmp_path, saved_base):
    payload["skills"][2]["base_override"] = saved_base
    workbook = TemplateWorkbook(catalog.template_path)
    path = tmp_path / "saved-specialization.xlsx"
    try:
        ExcelExporter(catalog)._write_character(workbook, build_draft(payload, catalog), None, tmp_path)
        workbook.save(path)
    finally:
        workbook.close()
    imported = import_investigator(path.read_bytes(), catalog)["draft"]
    row = next(item for item in imported["skills"] if item["template_slot"] == "F40")
    assert row["base_value"] == saved_base
    if saved_base != 15:
        assert row["base_override"] == saved_base
    else:
        assert "base_override" not in row
    rebuilt = next(item for item in build_draft(imported, catalog).skills if item.template_slot == "F40")
    assert (rebuilt.specialization, rebuilt.base_value, rebuilt.final_value) == ("冲锋枪", saved_base, saved_base + 13)


@pytest.mark.skipif(not shutil.which("libreoffice") and not shutil.which("soffice"), reason="LibreOffice Calc is required")
def test_named_subskills_keep_values_through_xlsx_recalculation_and_import(catalog, payload):
    draft = build_draft(payload, catalog)
    result = LinuxExcelExporter(catalog).export(draft)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(io.BytesIO(result.data), read_only=True, data_only=True)
    try:
        assert tuple(workbook.sheetnames) == catalog.sheet_names
        sheet = workbook["人物卡"]
        for slot in ("F38", "F39", "F40", "F41", "AB31"):
            skill = next(item for item in draft.skills if item.template_slot == slot)
            cells = catalog.skill_cell(slot)
            assert sheet[cells.specialization_cell].value == skill.specialization
            assert sheet[cells.base_cell].value == skill.base_value
            assert sheet[cells.total_cell].value == skill.final_value
    finally:
        workbook.close()
    imported = import_investigator(result.data, catalog)["draft"]
    rebuilt = {item.template_slot: item for item in build_draft(imported, catalog).skills}
    for original in draft.skills:
        recovered = rebuilt[original.template_slot]
        assert (recovered.name, recovered.specialization, recovered.base_value, recovered.final_value) == (
            original.name, original.specialization, original.base_value, original.final_value,
        )
