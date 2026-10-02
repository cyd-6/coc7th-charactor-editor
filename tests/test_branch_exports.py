"""Dynamic skill branches stay visible and round-trip through the original XLSX."""

from __future__ import annotations

import copy
import hashlib
import io
import shutil
import warnings
import zipfile

import pytest
from lxml import etree as ET
from openpyxl import load_workbook

from app import get_catalog
from coc7_card.branch_workbook import BRANCH_SCHEMA, branch_row, selection_title_row
from coc7_card.exporters.excel import ExcelExporter
from coc7_card.exporters.excel_linux import LinuxExcelExporter
from coc7_card.exporters.xlsx_template import TemplateWorkbook, tag, worksheet_paths
from coc7_card.importers.excel import ExcelImportError, import_investigator
from coc7_card.web import build_draft


@pytest.fixture(scope="module")
def catalog():
    return get_catalog()


def branch(index, name="射击", specialization="机枪", **values):
    slot = f"branch-{index:032x}"
    return {"key": slot, "template_slot": slot, "name": name,
            "specialization": specialization, **values}


@pytest.fixture(scope="module")
def payload():
    skills = [
        branch(1, specialization="机枪", base_override=0, occupation_points=12, interest_points=3,
               experience_points=4, extra_final=5),
        branch(2, specialization="重武器", interest_points=8),
        branch(3, specialization="喷射器", base_override=31, occupation_points=6),
        branch(4, "科学", "密码学", interest_points=9),
        branch(5, "科学", "海洋生物学", base_override=7, extra_final=2),
    ]
    return {
        "identity": {"name": "分支导出验证", "age": 30},
        "occupation": {"is_custom": True, "name": "分支验证职业", "credit_min": 0,
                       "credit_max": 99, "point_formula": "EDU*4"},
        "custom_skill_slots": [skills[0]["template_slot"], skills[2]["template_slot"], "F16"],
        "skills": skills,
        "experience": {"selection": "custom", "name": "分支经历", "skill_points": 10, "san_loss": 0},
    }


def write_fixture(catalog, payload, directory):
    workbook = TemplateWorkbook(catalog.template_path)
    path = directory / "branch-fixture.xlsx"
    try:
        ExcelExporter(catalog)._write_character(workbook, build_draft(payload, catalog), None, directory)
        workbook.save(path)
    finally:
        workbook.close()
    return path.read_bytes()


@pytest.fixture(scope="module")
def raw_workbook(catalog, payload, tmp_path_factory):
    return write_fixture(catalog, payload, tmp_path_factory.mktemp("branch-fixture"))


def worksheet(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    path = worksheet_paths(parts)["人物卡"]
    return parts, path, ET.fromstring(parts[path])


def replace_cell(data, address, *, value=None, formula=None):
    parts, path, root = worksheet(data)
    cell = next(node for node in root.iter(tag("c")) if node.get("r") == address)
    cell[:] = []
    cell.attrib.pop("t", None)
    if formula is not None:
        ET.SubElement(cell, tag("f")).text = formula
    elif isinstance(value, str):
        cell.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(cell, tag("is")), tag("t")).text = value
    elif value is not None:
        ET.SubElement(cell, tag("v")).text = str(value)
    parts[path] = ET.tostring(root)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return output.getvalue()


def api_payload(state):
    result = copy.deepcopy(state)
    if state["occupation_mode"] == "custom":
        custom = state["custom_occupation"]
        result["occupation"] = {**custom, "is_custom": True, "point_formula": "EDU*4"}
    return result


def assert_branch_values(original, recovered):
    actual = {skill.template_slot: skill for skill in recovered.skills}
    for skill in original.skills:
        other = actual[skill.template_slot]
        assert (other.name, other.specialization, other.base_value, other.occupation_points,
                other.interest_points, other.experience_points, other.extra_final, other.selected_occupation) == (
            skill.name, skill.specialization, skill.base_value, skill.occupation_points,
            skill.interest_points, skill.experience_points, skill.extra_final, skill.selected_occupation,
        )


def test_visible_branch_rows_preserve_template_structure(catalog, payload, raw_workbook):
    before = hashlib.sha256(catalog.template_path.read_bytes()).hexdigest()
    parts, _, root = worksheet(raw_workbook)
    source_parts, _, source = worksheet(catalog.template_path.read_bytes())
    assert tuple(worksheet_paths(parts)) == catalog.sheet_names
    merges = {entry.get("ref") for entry in root.find(tag("mergeCells"))}
    assert {entry.get("ref") for entry in source.find(tag("mergeCells"))} <= merges
    for name in source_parts:
        if name.endswith(".vml") or "/charts/" in name:
            assert parts[name] == source_parts[name]
    rows = {int(row.get("r")): row for row in root.find(tag("sheetData"))}
    cells = {cell.get("r"): cell for cell in root.iter(tag("c"))}
    assert "".join(cells["B153"].itertext()) == BRANCH_SCHEMA
    assert rows[153].get("hidden") == "1"
    for index, item in enumerate(payload["skills"]):
        row = branch_row(index)
        assert rows[row].get("hidden") != "1"
        assert rows[row + 1].get("hidden") == "1"
        assert "".join(cells[f"F{row}"].itertext()) == item["specialization"]
        assert "".join(cells[f"F{row + 1}"].itertext()) == item["template_slot"]
        assert cells[f"AG{row}"].findtext(tag("f")) == f"SUM(R{row}:AF{row})"
    formula = cells["J50"].findtext(tag("f"))
    for extra in ("U155:W163", "X155:Z163", "AA155:AC163"):
        assert extra in formula
    assert root.find(tag("dimension")).get("ref").endswith("169")
    assert hashlib.sha256(catalog.template_path.read_bytes()).hexdigest() == before


def test_custom_branch_ids_and_values_survive_import(catalog, payload, raw_workbook):
    state = import_investigator(raw_workbook, catalog)["draft"]
    assert len(state["skills"]) == len(catalog.skills) + 5
    rows = {skill["template_slot"]: skill for skill in state["skills"]}
    assert "base_override" not in rows[branch(2)["template_slot"]]
    assert rows[branch(1)["template_slot"]]["base_override"] == 0
    assert rows[branch(3)["template_slot"]]["base_override"] == 31
    assert rows[branch(5)["template_slot"]]["base_override"] == 7
    assert set(state["custom_skill_slots"]) == set(payload["custom_skill_slots"])
    assert_branch_values(build_draft(payload, catalog), build_draft(api_payload(state), catalog))


def test_free_and_group_choices_restore_exact_branch_ids(catalog, tmp_path):
    occupation = next(item for item in catalog.occupations if item.credit_min == 0 and item.free_choices > 0)
    skill = branch(30, "科学", "深海地质学", occupation_points=7)
    payload = {"identity": {"name": "目录分支验证"}, "occupation_id": occupation.occupation_id,
               "skills": [skill], "free_skill_choices": [skill["template_slot"]],
               "group_choices": {group.marker: list(group.candidates[:group.required_count])
                                 for group in occupation.choice_groups}}
    state = import_investigator(write_fixture(catalog, payload, tmp_path), catalog)["draft"]
    assert state["free_skill_choices"] == payload["free_skill_choices"]
    assert state["group_choices"] == payload["group_choices"]
    assert_branch_values(build_draft(payload, catalog), build_draft(state, catalog))


def test_fixed_occupation_specialization_reuses_branch_after_import(catalog, tmp_path):
    occupation = next(item for item in catalog.occupations if item.name == "建筑师")
    skill = branch(40, "科学", "数学", occupation_points=7)
    payload = {"identity": {"name": "固定分支验证"}, "occupation_id": occupation.occupation_id,
               "skills": [skill]}
    original = build_draft(payload, catalog)
    assert next(item for item in original.skills if item.template_slot == skill["template_slot"]).selected_occupation
    state = import_investigator(write_fixture(catalog, payload, tmp_path), catalog)["draft"]
    assert next(item for item in state["skills"] if item["template_slot"] == skill["template_slot"])["selected_occupation"]
    assert not next(item for item in state["skills"] if item["template_slot"] == "AB31")["selected_occupation"]
    assert_branch_values(original, build_draft(state, catalog))


@pytest.mark.parametrize("address,value,formula,error", [
    ("F158", "branch-00000000000000000000000000000001", None, "重复分支标识"),
    ("U155", None, "1+1", "已保存的非负整数"),
    ("R155", "无法识别", None, "已保存的非负整数"),
    ("B153", "COC7_BRANCHES_V99", None, "版本标记"),
    ("AP155", "", None, "本职标志"),
    ("AP153", 4, None, "数量与分支标识"),
    ("AP166", 2, None, "记录数量不完整"),
    ("F155", "手枪", None, "与已有分项重名"),
])
def test_broken_branch_metadata_is_not_silently_dropped(catalog, raw_workbook, address, value, formula, error):
    changed = replace_cell(raw_workbook, address, value=value, formula=formula)
    with pytest.raises(ExcelImportError, match=error):
        import_investigator(changed, catalog)


def test_branch_specialization_is_literal_text(catalog, tmp_path):
    payload = {"identity": {"name": "字面文本验证"}, "occupation_id": catalog.occupations[0].occupation_id,
               "skills": [branch(99, "科学", "=2+2")]}
    data = write_fixture(catalog, payload, tmp_path)
    _, _, root = worksheet(data)
    cell = next(node for node in root.iter(tag("c")) if node.get("r") == "F155")
    assert cell.get("t") == "inlineStr" and cell.find(tag("f")) is None
    assert import_investigator(data, catalog)["draft"]["skills"][-1]["specialization"] == "=2+2"


@pytest.mark.skipif(not shutil.which("libreoffice") and not shutil.which("soffice"), reason="LibreOffice Calc is required")
def test_real_recalculation_branch_caches_budgets_and_roundtrip(catalog, payload):
    original_hash = hashlib.sha256(catalog.template_path.read_bytes()).hexdigest()
    draft = build_draft(payload, catalog)
    exported = LinuxExcelExporter(catalog).export(draft)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(io.BytesIO(exported.data), read_only=True, data_only=True)
    try:
        assert tuple(workbook.sheetnames) == catalog.sheet_names
        sheet = workbook["人物卡"]
        for index, skill in enumerate(draft.skills[-5:]):
            row = branch_row(index)
            assert sheet[f"AG{row}"].value == skill.final_value
            assert sheet[f"AJ{row}"].value == skill.hard_value
            assert sheet[f"AM{row}"].value == skill.extreme_value
        budget = sheet["J50"].value
        assert "剩余职业点=182" in budget
        assert "剩余兴趣点=80" in budget
        assert "剩余经历包点=6" in budget
    finally:
        workbook.close()
    state = import_investigator(exported.data, catalog)["draft"]
    assert set(state["custom_skill_slots"]) == set(payload["custom_skill_slots"])
    assert_branch_values(draft, build_draft(api_payload(state), catalog))
    assert hashlib.sha256(catalog.template_path.read_bytes()).hexdigest() == original_hash
