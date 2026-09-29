from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import shutil
import sys
import warnings
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from lxml import etree as ET
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from app import FONT_PATH, app, get_catalog
from coc7_card.exporters.excel import ExcelExporter, ExcelExportError
from coc7_card.exporters.excel_linux import LinuxExcelExporter
from coc7_card.exporters.pdf import PdfExporter, PdfExportError
from coc7_card.exporters.xlsx_template import TemplateWorkbook, tag
from coc7_card.web import build_draft
from scripts.recalculate_template import recalculate


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as connection:
        yield connection


@pytest.fixture
def payload():
    catalog = get_catalog()
    occupation = next(item for item in catalog.occupations if item.credit_min == 0)
    return {
        "identity": {
            "name": "林知秋", "player": "容器测试", "age": 32,
            "story_year": 1920, "era": "1920s", "current_date": "1920-04-12",
            "gender": "女", "birthplace": "上海", "residence": "阿卡姆",
        },
        "attributes": {"STR": 60, "CON": 65, "SIZ": 55, "DEX": 70, "APP": 50, "INT": 75, "POW": 60, "EDU": 70, "Luck": 55},
        "occupation_id": occupation.occupation_id,
        "background": {"appearance": "短发，灰色风衣", "beliefs": "寻找事实", "personal_story": "调查一封来自故友的信。"},
        "weapons": [{"name": "小刀", "skill": "斗殴", "damage": "1D4+DB", "range": "接触", "attacks": "1"}],
        "inventory": [{"name": "调查笔记", "status": "携带", "location": "外套口袋", "backpack_slot": "1"}],
    }


def portrait_bytes():
    buffer = io.BytesIO()
    Image.new("RGB", (80, 120), "#4c7084").save(buffer, format="PNG")
    return buffer.getvalue()


def export(client, kind, payload, portrait=False):
    files = {"portrait": ("portrait.png", portrait_bytes(), "image/png")} if portrait else None
    return client.post(f"/api/export/{kind}", data={"draft_json": json.dumps(payload, ensure_ascii=False)}, files=files)


def read_book(data, data_only):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return load_workbook(io.BytesIO(data), data_only=data_only)


def test_web_catalog_calculation_and_validation(client, payload):
    assert client.get("/").status_code == 200
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/assets/fonts/NotoSansSC-Regular.ttf").status_code == 200
    assert client.get("/api/health").json()["status"] == "ok"
    catalog = client.get("/api/bootstrap").json()
    assert catalog["meta"]["sheet_count"] == 14
    assert catalog["meta"]["occupation_count"] > 200
    result = client.post("/api/calculate", json={"attributes": payload["attributes"], "age": 32, "occupation_formula": "EDU*4"}).json()["derived"]
    assert (result["hp"], result["mp"], result["san"], result["dodge"]) == (12, 12, 60, 35)
    assert (result["occupation_points"], result["interest_points"]) == (280, 150)
    assert client.post("/api/validate", json=payload).json()["validation"]["can_export"]


def test_pdf_matches_two_reference_pages_with_portrait(client, payload):
    response = export(client, "pdf", payload, portrait=True)
    assert response.status_code == 200, response.text
    result = response.json()
    reader = PdfReader(io.BytesIO(base64.b64decode(result["pdf_base64"])))
    assert len(reader.pages) == result["page_count"] == 2
    assert not reader.get_fields()
    text = "".join(page.extract_text() for page in reader.pages)
    assert "林知秋" in text and "调查笔记" in text
    reference = PdfReader(Path(__file__).resolve().parents[1] / "assets/templates/1920sCha.pdf")
    for page, original in zip(reader.pages, reference.pages, strict=True):
        assert tuple(float(value) for value in page.mediabox) == pytest.approx(tuple(float(value) for value in original.mediabox), abs=0.01)
        assert tuple(float(value) for value in page.cropbox) == pytest.approx(tuple(float(value) for value in original.cropbox), abs=0.01)
    assert len(result["previews"]) == 2
    for preview in result["previews"]:
        Image.open(io.BytesIO(base64.b64decode(preview))).verify()


requires_calc = pytest.mark.skipif(not shutil.which("libreoffice") and not shutil.which("soffice"), reason="LibreOffice Calc is required")


@requires_calc
@pytest.mark.parametrize("year,currency,custom,portrait", [(1920, "USD", False, False), (1925, "GBP", False, False), (2026, "EUR", True, True)])
def test_excel_preserves_template_and_recalculates(client, payload, year, currency, custom, portrait):
    catalog = get_catalog()
    before = hashlib.sha256(catalog.template_path.read_bytes()).hexdigest()
    payload["identity"].update(story_year=year, era="现代" if year == 2026 else "1920s")
    payload["assets"] = {"currency_code": currency, "usd_cash": "123.45", "usd_assets": "0", "usd_spending": ""}
    credit = next(skill for skill in catalog.skills if skill.name.rstrip("：:") == "信用评级")
    payload["skills"] = [{"template_slot": credit.template_slot, "interest_points": 30}]
    if custom:
        # The name is literal text, even if it starts with an Excel formula prefix.
        payload["identity"]["name"] = "=1+1"
        payload["occupation"] = {"is_custom": True, "name": "民俗研究员", "credit_min": 0, "credit_max": 99, "point_formula": "EDU*2+INT*2"}
        payload["custom_skill_slots"] = [skill.template_slot for skill in catalog.skills[:8]]
    else:
        # The chosen catalog occupation must permit the allocated credit rating.
        occupation = next(item for item in catalog.occupations if item.credit_min <= 30 <= item.credit_max)
        payload["occupation_id"] = occupation.occupation_id
    response = export(client, "excel", payload, portrait=portrait)
    assert response.status_code == 200, response.text
    assert response.headers["X-Workbook-Sheet-Count"] == "14"
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    formulas, values = read_book(response.content, False), read_book(response.content, True)
    try:
        assert tuple(formulas.sheetnames) == catalog.sheet_names
        assert values["人物卡"]["E3"].value == payload["identity"]["name"]
        assert formulas["人物卡"]["E3"].data_type == "s"
        assert values["人物卡"]["U3"].value == 60
        assert values["人物卡"]["E10"].value == 12
        assert values["人物卡"]["R26"].value == 30
        assert values["人物卡"]["F79"].value == "调查笔记"
        assert values["人物卡"]["B53"].value == "小刀"
        assert values["人物卡"]["AA61"].value == "短发，灰色风衣"
        rate = values["货币汇率"]["K5"].value
        quote = next(item for item in client.get("/api/bootstrap").json()["currencies"][str(year)] if item["code"] == currency)
        assert rate == pytest.approx(float(quote["rate"]), rel=1e-6)
        assert values["人物卡"]["O62"].value == pytest.approx(round(123.45 * rate, 2), abs=0.01)
        assert values["人物卡"]["L62"].value == 0
        assert values["人物卡"]["I62"].value == pytest.approx(round(10 * (20 if year == 2026 else 1) * rate, 2), abs=0.01)
        assert values["货币汇率"]["K11"].value == 0
        assert formulas["货币汇率"]["K9"].data_type == "f"
        assert formulas["更新说明"]["P3"].value is None
        assert values["附表"]["J6"].value != 0  # The source template caches zero here.
        assert sum(values["附表"][f"{column}8"].value for column in "GHIJKL") == pytest.approx(1)
        if custom:
            assert values["人物卡"]["E5"].value == "民俗研究员"
            assert values["职业列表"]["F3"].value == 290
            assert len(formulas["人物卡"]._images) == 1
        for sheet in values:
            assert not [(cell.coordinate, cell.value) for row in sheet for cell in row if cell.data_type == "e"], sheet.title
    finally:
        formulas.close()
        values.close()
    # OOXML features unsupported by openpyxl and altered by LibreOffice must
    # survive byte-for-byte where untouched, including chart and VML parts.
    with zipfile.ZipFile(catalog.template_path) as source, zipfile.ZipFile(io.BytesIO(response.content)) as output:
        for name in source.namelist():
            if name.startswith("xl/charts/") or name.endswith(".vml"):
                assert output.read(name) == source.read(name)
        original_xml = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
        exported_xml = ET.fromstring(output.read("xl/worksheets/sheet1.xml"))
        for element in ["extLst", "dataValidations", "mergeCells"]:
            assert ET.tostring(original_xml.find(tag(element)), method="c14n") == ET.tostring(exported_xml.find(tag(element)), method="c14n")
        assert "xl/calcChain.xml" not in output.namelist()
    assert hashlib.sha256(catalog.template_path.read_bytes()).hexdigest() == before


def test_invalid_draft_never_exports(client, payload):
    payload["identity"]["name"] = ""
    for kind in ["pdf", "excel"]:
        response = export(client, kind, payload)
        assert response.status_code == 422
        assert "姓名不能为空" in response.json()["detail"]


@pytest.mark.parametrize("kind", ["windows", "linux", "pdf"])
def test_public_export_rejects_missing_occupation(payload, kind):
    catalog = get_catalog()
    draft = build_draft(payload, catalog)
    draft.occupation = None
    if kind == "pdf":
        exporter = PdfExporter(FONT_PATH)
        error = PdfExportError
    else:
        exporter = (ExcelExporter if kind == "windows" else LinuxExcelExporter)(catalog)
        error = ExcelExportError
    with pytest.raises(error, match="必须选择或创建一个职业"):
        exporter.export(draft)


@requires_calc
@pytest.mark.parametrize("sheet_name,address,value,message", [
    ("简化卡 骰娘导入", "T29", '=IF(W29=0,"",人物卡!AB41)', "模板兼容性修复未写入"),
    ("更新说明", "P3", "旧模板版本说明", "仍含模板版本说明"),
    ("附表", "AB61", "=1/0", "公式错误"),
])
def test_excel_rejects_unprepared_template(payload, tmp_path, sheet_name, address, value, message):
    catalog = get_catalog()
    template = tmp_path / "unprepared.xlsx"
    book = TemplateWorkbook(catalog.template_path)
    try:
        target = book.Worksheets(sheet_name).Range(address)
        if value.startswith("="):
            target.Formula = value
        else:
            target.Value2 = value
        book.save(template)
    finally:
        book.close()
    # Keep valid schema metadata and a matching source hash: rejection must
    # come from the output checks, without repairing this altered template.
    template_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    altered_catalog = replace(catalog, template_path=template, source_sha256=template_hash)
    with pytest.raises(ExcelExportError, match=message):
        LinuxExcelExporter(altered_catalog).export(build_draft(payload, altered_catalog))
    assert hashlib.sha256(template.read_bytes()).hexdigest() == template_hash


@requires_calc
@pytest.mark.parametrize("package", [None, "custom", "战场经历包"])
def test_excel_separates_growth_from_experience_budget(client, payload, package):
    catalog = get_catalog()
    budget = 0
    if package == "custom":
        budget = 40
        payload["experience"] = {"selection": "custom", "name": "档案调查", "skill_points": budget, "san_loss": 0}
    elif package:
        budget = next(p["skill_points"] for p in catalog.experience_packages if p["name"] == package)
        payload["experience"] = {"selection": package, "san_loss": 0}
    examples = [("F16", 2, 12, 17), ("AB16", 3, 7, 11), ("F20", 0, 5, 3)]
    payload["skills"] = [
        {"template_slot": slot, "interest_points": interest, "extra_final": growth,
         "experience_points": experience if package else 0}
        for slot, interest, growth, experience in examples
    ]
    response = export(client, "excel", payload)
    assert response.status_code == 200, response.text
    values, formulas = read_book(response.content, True), read_book(response.content, False)
    try:
        card = values["人物卡"]
        assert card["L15"].value.replace("\n", "") == "成长点数"
        assert card["M15"].value.replace("\n", "") == "经历包点"
        assert card["AH15"].value == card["L15"].value
        assert card["AI15"].value == card["M15"].value
        for slot, interest, growth, experience in examples:
            mapping = catalog.skill_cell(slot)
            row = int(slot.lstrip("FAB"))
            hard, extreme = ("T", "V") if slot.startswith("F") else ("AP", "AR")
            final = 5 + interest + growth + (experience if package else 0)
            assert card[mapping.extra_cell].value == growth
            assert card[mapping.experience_cell].value == (experience if package else 0)
            assert card[mapping.total_cell].value == final
            assert card[f"{hard}{row}"].value == final // 2
            assert card[f"{extreme}{row}"].value == final // 5
            assert formulas["人物卡"][mapping.total_cell].data_type == "f"
        remaining = budget - (31 if package else 0)
        assert values["附表"]["AC27"].value == remaining
        if package:
            assert f"剩余经历包点={remaining}" in card["J50"].value
        else:
            assert "剩余经历包点" not in card["J50"].value
        for sheet in values:
            assert not [cell.coordinate for row in sheet for cell in row if cell.data_type == "e"], sheet.title
        # The independent package inputs must really be writable, including
        # skill rows whose previous total used explicit addition, not SUM.
        for mapping in catalog.skill_cells:
            assert not any(mapping.experience_cell in region for region in formulas["人物卡"].merged_cells)
    finally:
        values.close()
        formulas.close()


@requires_calc
def test_editing_exported_point_columns_recalculates_independently(client, payload, tmp_path):
    payload["experience"] = {"selection": "custom", "name": "独立换算", "skill_points": 40}
    payload["skills"] = [{"template_slot": "F16", "extra_final": 12, "experience_points": 17},
                         {"template_slot": "AB16", "extra_final": 7, "experience_points": 11}]
    response = export(client, "excel", payload)
    assert response.status_code == 200, response.text
    baseline = tmp_path / "export.xlsx"
    baseline.write_bytes(response.content)
    # Editing growth changes success rates, but never the package budget.
    # Editing package points changes both, even when other budgets are spent.
    for cells, expected_left, expected_right, remaining in [
        ({"L16": 20, "AH16": 9, "N18": 280, "P19": 150}, 42, 25, 12),
        ({"L16": 20, "AH16": 9, "M16": 18, "AI16": 12}, 43, 26, 10),
    ]:
        book = TemplateWorkbook(baseline)
        edited, calculated = tmp_path / "edited.xlsx", tmp_path / "calculated.xlsx"
        try:
            for address, value in cells.items():
                book.Worksheets("人物卡").Range(address).Value2 = value
            book.save(edited)
        finally:
            book.close()
        recalculate(edited, calculated)
        values = read_book(calculated.read_bytes(), True)
        try:
            card = values["人物卡"]
            assert card["R16"].value == expected_left
            assert card["AN16"].value == expected_right
            assert values["附表"]["AC27"].value == remaining
            assert f"剩余经历包点={remaining}" in card["J50"].value
        finally:
            values.close()


def test_missing_libreoffice_has_actionable_error(client, payload, monkeypatch):
    monkeypatch.setattr("coc7_card.exporters.excel_linux.shutil.which", lambda _name: None)
    response = export(client, "excel", payload)
    assert response.status_code == 422
    assert "LibreOffice Calc" in response.json()["detail"]


def test_recalculation_timeout_stops_process(tmp_path):
    runner = tmp_path / "slow-office"
    pid_file = tmp_path / "process.pid"
    runner.write_text(f"#!{sys.executable}\nimport os, time\nfrom pathlib import Path\nPath({str(pid_file)!r}).write_text(str(os.getpid()))\ntime.sleep(60)\n")
    runner.chmod(0o700)
    with pytest.raises(ExcelExportError, match="超过 1 秒"):
        LinuxExcelExporter._recalculate(str(runner), tmp_path / "input.xlsx", tmp_path / "output", tmp_path / "profile", 1)
    with pytest.raises(ProcessLookupError):
        os.kill(int(pid_file.read_text()), 0)


@requires_calc
def test_parallel_exports_do_not_mix_characters(payload):
    def generate(name):
        own = {**payload, "identity": {**payload["identity"], "name": name}}
        with TestClient(app) as connection:
            response = export(connection, "excel", own)
            assert response.status_code == 200, response.text
            book = read_book(response.content, True)
            try:
                return book["人物卡"]["E3"].value
            finally:
                book.close()
    with ThreadPoolExecutor(max_workers=2) as executor:
        assert list(executor.map(generate, ["调查员甲", "调查员乙"])) == ["调查员甲", "调查员乙"]
