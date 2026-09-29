"""Regression coverage for bounded-memory catalog loading and XLSX checks."""
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from lxml import etree as ET
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from app import TEMPLATE_PATH
from coc7_card.catalog import TemplateCatalog
from coc7_card.exporters.excel import ExcelExporter, ExcelExportError
from coc7_card.exporters.xlsx_template import MAIN, REL, PACKAGE, tag


def test_catalog_streams_without_random_cell_reads(monkeypatch):
    import coc7_card.catalog as module
    original = module.load_workbook
    def load(*args, **kwargs):
        assert kwargs['read_only'] is True
        return original(*args, **kwargs)
    def random_access(*args, **kwargs):
        pytest.fail('Read-only random cell access reparses a worksheet')
    monkeypatch.setattr(module, 'load_workbook', load)
    monkeypatch.setattr(ReadOnlyWorksheet, 'cell', random_access)
    catalog = TemplateCatalog.load(TEMPLATE_PATH)
    assert (len(catalog.occupations), len(catalog.skills), len(catalog.weapons), len(catalog.inventory)) == (229, 67, 104, 430)
    assert catalog.skill_cell('AB16').experience_cell == 'AI16'
    assert catalog.skill_cell('F16').experience_cell == 'M16'


@pytest.fixture
def verification_package():
    names = ('更新说明', '简化卡 骰娘导入', '附表', '人物卡')
    workbook = ET.Element(tag('workbook'), nsmap={None: MAIN, 'r': REL})
    sheets = ET.SubElement(workbook, tag('sheets'))
    relations = ET.Element(f'{{{PACKAGE}}}Relationships')
    roots = {}
    for index, name in enumerate(names, 1):
        ET.SubElement(sheets, tag('sheet'), name=name, sheetId=str(index), attrib={f'{{{REL}}}id':f'rId{index}'})
        ET.SubElement(relations, f'{{{PACKAGE}}}Relationship', Id=f'rId{index}', Target=f'worksheets/sheet{index}.xml')
        roots[name] = ET.Element(tag('worksheet'), nsmap={None: MAIN})
        ET.SubElement(roots[name], tag('sheetData'))
    def cell(sheet, address, formula=None, value='0', kind=None, **attributes):
        data = roots[sheet].find(tag('sheetData'))
        row = ET.SubElement(data, tag('row'), r=''.join(filter(str.isdigit, address)))
        node = ET.SubElement(row, tag('c'), r=address)
        if kind:
            node.set('t', kind)
        if formula is not None:
            ET.SubElement(node, tag('f'), **attributes).text = formula
        ET.SubElement(node, tag('v')).text = value
        return node
    cell('简化卡 骰娘导入', 'T28', 'IF(W28=0,"",人物卡!AB41)', t='shared', si='0', ref='T28:T29')
    cell('简化卡 骰娘导入', 'T29', '', t='shared', si='0')
    cell('简化卡 骰娘导入', 'T34', 'IF(W34=0,"",人物卡!AB47)')
    for address in ('AT15', 'AT16', 'AT17'):
        cell('附表', address, '附表!$V$226')
    strings = ET.Element(tag('sst'))
    ET.SubElement(ET.SubElement(strings, tag('si')), tag('t')).text = '旧说明'
    def save(path):
        with ZipFile(path, 'w') as archive:
            archive.writestr('xl/workbook.xml', ET.tostring(workbook))
            archive.writestr('xl/_rels/workbook.xml.rels', ET.tostring(relations))
            archive.writestr('xl/sharedStrings.xml', ET.tostring(strings))
            for index, name in enumerate(names, 1):
                archive.writestr(f'xl/worksheets/sheet{index}.xml', ET.tostring(roots[name]))
        return path
    return SimpleNamespace(names=names, roots=roots, cell=cell, save=save)


def test_stream_verifier_accepts_shared_formulas(verification_package, tmp_path):
    fixture = verification_package
    assert ExcelExporter(SimpleNamespace(sheet_names=fixture.names))._verify_export(fixture.save(tmp_path/'valid.xlsx')) == 4


def test_stream_verifier_preserves_only_the_loaded_template_revision(verification_package, tmp_path):
    fixture = verification_package
    fixture.cell('更新说明','P3',None,'0','s')
    exporter = ExcelExporter(SimpleNamespace(sheet_names=fixture.names,template_revision_note='旧说明'))
    path = fixture.save(tmp_path/'matching-note.xlsx')
    assert exporter._verify_export(path) == 4
    changed = ExcelExporter(SimpleNamespace(sheet_names=fixture.names,template_revision_note='新的已保存说明'))
    with pytest.raises(ExcelExportError,match='版本说明与当前模板不一致'):
        changed._verify_export(path)
    data = fixture.roots['更新说明'].find(tag('sheetData'))
    data.remove(data[0])
    with pytest.raises(ExcelExportError,match='缺少当前模板的版本说明'):
        exporter._verify_export(fixture.save(tmp_path/'missing-note.xlsx'))


@pytest.mark.parametrize('sheet,address,formula,value,kind,message', [
    ('人物卡', 'A1', '1/0', '#DIV/0!', 'e', '公式错误'),
    ('人物卡', 'A1', '#REF!+1', '0', None, '断裂引用'),
    ('人物卡', 'A1', None, '#VALUE!', 'e', '公式错误'),
    ('更新说明', 'P3', None, '0', 's', '仍含模板版本说明'),
    ('更新说明', 'P3', '""', '', 'str', '仍含模板版本说明'),
])
def test_stream_verifier_rejects_bad_outputs(verification_package, tmp_path, sheet, address, formula, value, kind, message):
    fixture = verification_package
    fixture.cell(sheet, address, formula, value, kind)
    with pytest.raises(ExcelExportError, match=message):
        ExcelExporter(SimpleNamespace(sheet_names=fixture.names))._verify_export(fixture.save(tmp_path/'bad.xlsx'))


def test_stream_verifier_requires_compatibility_cells_and_sheet_order(verification_package, tmp_path):
    fixture = verification_package
    path = fixture.save(tmp_path/'valid.xlsx')
    with pytest.raises(ExcelExportError, match='数量或顺序'):
        ExcelExporter(SimpleNamespace(sheet_names=fixture.names[::-1]))._verify_export(path)
    data = fixture.roots['附表'].find(tag('sheetData'))
    data.remove(data[-1])
    with pytest.raises(ExcelExportError, match='附表!AT17'):
        ExcelExporter(SimpleNamespace(sheet_names=fixture.names))._verify_export(fixture.save(tmp_path/'missing.xlsx'))
