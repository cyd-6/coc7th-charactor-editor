"""Source reconciliation and behavior at historical-unit/missing-data boundaries."""
import base64
import io
import json
import os
import shutil
import subprocess
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app import app, get_catalog
from coc7_card.currency import currency_options, format_amount, load_exchange_data
from coc7_card.exporters.xlsx_template import TemplateWorkbook
from scripts.recalculate_template import recalculate

ROOT = Path(__file__).resolve().parents[1]
DATA = load_exchange_data()


def quote(year, code):
    return next(q for q in DATA['quotes'] if q['year']==year and q['code']==code)


def payload(year=2026, code='GBP'):
    catalog=get_catalog()
    return {'version':3, 'identity':{'name':'汇率校验','age':32,'story_year':year,'era':'1920s'},
            'attributes':dict(STR=60,CON=65,SIZ=55,DEX=70,APP=50,INT=75,POW=60,EDU=70,Luck=55),
            'occupation_id':next(o.occupation_id for o in catalog.occupations if o.credit_min==0),
            'assets':{'currency_code':code,'usd_cash':'123.45','usd_assets':'0','usd_spending':''}}


def test_every_source_cell_is_preserved_once_or_an_explicit_proxy():
    source=load_workbook(ROOT/'assets/rates/sources'/DATA['source_file'],read_only=False,data_only=True)
    try:
        seen=set()
        for q in DATA['quotes']:
            for location in q['source_cells']:
                sheet,cell=location.split('!')
                value=source[sheet][cell].value
                if q['available']:
                    assert Decimal(str(value))==pytest.approx(Decimal(q['rate']),rel=Decimal('1e-15'))
                else:
                    assert value is None
                    assert q['missing_reason']
                assert location not in seen
                seen.add(location)
        assert len(seen)==107*12
        assert len(DATA['quotes'])==1121
        assert sum(not q['available'] for q in DATA['quotes'])==60
        for q in DATA['quotes']:
            if q['available'] and q['code']!='USD':
                assert q['source_url'].startswith('https://')
                assert q['source'] and q['raw_cell'] and q['conversion']
    finally: source.close()


def test_all_years_and_source_metadata_are_served():
    with TestClient(app) as client:
        response=client.get('/api/bootstrap')
        assert response.status_code==200
        result=response.json()
        assert set(result['currencies'])=={str(y) for y in range(1920,2027)}
        assert result['currency_metadata']['source_sha256']==DATA['source_sha256']
        assert sum(len(v) for v in result['currencies'].values())==1061
        assert sum(len(v) for v in result['currency_missing'].values())==60
        assert not {'AUD','HKD','KRW'} & {q['code'] for q in result['currencies']['2026']}


@pytest.mark.parametrize('year,code',[(1919,'USD'),(2027,'USD'),(1923,'DE_TRANSITION_1923'),(1920,'RUB_EARLY'),
                                    (1998,'EUR'),(2026,'AUD'),(1954,'CNY'),(1959,'FRF'),(1997,'RUB'),(2026,'')])
def test_invalid_selection_is_never_silently_converted_to_usd(year,code):
    with TestClient(app) as client:
        response=client.post('/api/validate',json=payload(year,code))
        assert response.status_code==422
        assert '年份' in response.json()['detail'] or '币种' in response.json()['detail']


def test_legacy_draft_aliases_only_identical_units():
    with TestClient(app) as client:
        draft=payload(1929,'FRF');draft['version']=2
        assert client.post('/api/validate',json=draft).status_code==200
        draft['assets']['currency_code']='SILVER_CNY'
        assert client.post('/api/validate',json=draft).status_code==422


def test_story_year_takes_precedence_over_stale_asset_year():
    draft=payload(2026,'GBP');draft['assets']['exchange_year']='old-draft-value'
    with TestClient(app) as client:
        assert client.post('/api/validate',json=draft).status_code==200


def test_unit_transitions_and_ytd_are_explicit():
    for year,code in [(1921,'TAEL'),(1922,'CN_SILVER_OLD'),(1954,'CNY_OLD'),(1955,'CNY'),
                      (1959,'FRF_OLD'),(1960,'FRF'),(1997,'RUR'),(1998,'RUB'),(1999,'EUR')]:
        assert quote(year,code)['available']
    assert '8 月' in quote(2026,'GBP')['warning']
    assert quote(2026,'GBP')['date']=='2026-01-01—2026-08-31'
    assert len([q for q in currency_options(DATA['quotes'],1999) if '欧元' in q['name']])==1
    assert not any(q['code']=='EUR' for q in currency_options(DATA['quotes'],1998))


def test_template_contains_the_same_dataset_and_yearly_dropdowns():
    book=load_workbook(ROOT/'assets/templates/COC7空白卡CY26.2.xlsx',data_only=False)
    try:
        assert len(book.sheetnames)==14
        assert book.defined_names['COC7_FX_DATA_SHA256'].attr_text.strip('"')==DATA['data_sha256']
        sheet=book['货币汇率']
        assert not sheet.row_dimensions[2].hidden
        for row,q in enumerate(DATA['quotes'],36):
            assert sheet.cell(row,1).value==q['year']
            assert sheet.cell(row,2).value==q['code']
            assert sheet.cell(row,3).value==q['name']
            assert sheet.cell(row,7).value==q['source']
            assert (sheet.cell(row,8).value or '')==q['source_url']
            rate=sheet.cell(row,5).value
            if q['available']: assert rate==pytest.approx(float(q['rate']),rel=1e-14)
            else: assert rate is None
        for year in range(1920,2027):
            destinations=list(book.defined_names[f'FX_Y{year}'].destinations)
            assert len(destinations)==1 and destinations[0][0]=='货币汇率'
            actual=[c.value for row in sheet[destinations[0][1]] for c in row]
            assert actual==[q['name'] for q in currency_options(DATA['quotes'],year)]
        validations={str(v.sqref):v for v in sheet.data_validations.dataValidation}
        assert validations['K2'].formula1=='INDIRECT("FX_Y"&$K$1)'
        assert validations['K1'].formula1=='1920'
        assert validations['K1'].formula2=='2026'
    finally: book.close()


def test_decimal_rounding_is_identical_in_browser_and_server():
    cases=[['1.005','1'],['0','123.456'],['1e-3','5'],['1000000000000000','1.234567891234567'],
           *[['123.45',q['rate']] for q in DATA['quotes'] if q['available']]]
    command="const f=require('./static/currency.js');let s='';process.stdin.on('data',x=>s+=x);process.stdin.on('end',()=>process.stdout.write(JSON.stringify(JSON.parse(s).map(([a,r])=>f.formatAmount(a,r)))));"
    node=shutil.which('node') or os.environ.get('COC7_NODE')
    if not node: pytest.skip('Node.js is required to verify the browser decimal helper')
    run=subprocess.run([node,'-e',command],input=json.dumps(cases),text=True,capture_output=True,cwd=ROOT,check=True)
    assert json.loads(run.stdout)==[format_amount(a,r) for a,r in cases]


@pytest.mark.skipif(not shutil.which('libreoffice'),reason='LibreOffice required')
def test_exported_excel_can_change_year_and_currency_independently(tmp_path):
    with TestClient(app) as client:
        response=client.post('/api/export/excel',data={'draft_json':json.dumps(payload(1920,'USD'))})
        assert response.status_code==200,response.text
    baseline=tmp_path/'baseline.xlsx';baseline.write_bytes(response.content)
    cases=[(1920,'TAEL'),(1954,'CNY_OLD'),(1955,'CNY'),(1959,'FRF_OLD'),(1960,'FRF'),
           (1997,'RUR'),(1998,'RUB'),(1999,'EUR'),(2026,'GBP'),(1923,'DE_TRANSITION_1923')]
    for year,code in cases:
        q=quote(year,code)
        # Simulate editing Excel inputs only, without calling the exporter again.
        book=TemplateWorkbook(baseline)
        try:
            fx=book.Worksheets('货币汇率')
            fx.Range('K1').Value2=year
            fx.Range('K2').Value2=q['name']
            incoming=tmp_path/f'input-{year}.xlsx';output=tmp_path/f'result-{year}.xlsx'
            book.save(incoming)
        finally: book.close()
        recalculate(incoming,output)
        values=load_workbook(output,data_only=True)
        try:
            rate=values['货币汇率']['K5'].value
            if q['available']:
                assert rate==pytest.approx(float(q['rate']),rel=1e-14)
                assert values['人物卡']['O62'].value==pytest.approx(float(format_amount('123.45',q['rate']).replace(',','')),abs=.001)
                assert values['人物卡']['S62'].value==q['name']
                assert values['货币汇率']['K7'].value==q['source']
                assert values['货币汇率']['K6'].value==q['date']
            else:
                assert rate in ('',None)
                assert values['人物卡']['O62'].value in ('',None)
                assert '缺少报价' in values['货币汇率']['K8'].value
            assert values['货币汇率']['K10'].value==123.45
        finally: values.close()


def test_pdf_records_quote_period_and_basis():
    with TestClient(app) as client:
        response=client.post('/api/export/pdf',data={'draft_json':json.dumps(payload())})
        assert response.status_code==200,response.text
        data=response.json()
        reader=PdfReader(io.BytesIO(base64.b64decode(data['pdf_base64'])))
        text=''.join(p.extract_text() for p in reader.pages)
        assert '2026-08-31' in text
        assert format_amount('123.45',quote(2026,'GBP')['rate']) in text
        assert 'FRED' in text and '均值' in text
