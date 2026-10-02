"""Windows launcher and isolated Excel automation; native tests are opt-in."""
import hashlib
import io
import json
import os
import socket
import sys
import threading
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import app, get_catalog
from coc7_card import launcher
from coc7_card.exporters.excel import ExcelExporter, ExcelExportError, _set_cell_value
from coc7_card.importers.excel import import_investigator
from coc7_card.web import build_draft


def sample_payload(name='Windows测试', year=1920, code='USD'):
    catalog = get_catalog()
    return {'identity': {'name': name, 'player': '00123', 'age': 32, 'story_year': year},
            'attributes': dict(STR=60,CON=65,SIZ=55,DEX=70,APP=50,INT=75,POW=60,EDU=70,Luck=55),
            'occupation_id': next(o.occupation_id for o in catalog.occupations if o.credit_min==0),
            'assets': {'currency_code':code, 'usd_cash':'123.45', 'usd_assets':'0'},
            'background': {'appearance':'=1+1', 'personal_story':'验证 Windows 导出。'}}


def test_launcher_reserves_a_free_port_when_preferred_is_busy():
    with launcher._new_listener_socket() as occupied:
        occupied.bind(('127.0.0.1', 0))
        occupied.listen(1)
        preferred = occupied.getsockname()[1]
        port, listener = launcher._reserve_listener(preferred_port=preferred, search_limit=1)
        try:
            assert port != preferred
            assert listener.getsockname() == ('127.0.0.1', port)
            with socket.socket() as other, pytest.raises(OSError):
                other.bind(('127.0.0.1', port))
        finally:
            listener.close()


@pytest.mark.parametrize('value', ['=1+1', '+SUM(A1)', '00123', "'前缀", '文字', '', 0, 12.5])
def test_com_values_preserve_text_but_ooxml_does_not_gain_a_quote(value):
    native = SimpleNamespace(_oleobj_=object())
    xml = SimpleNamespace()
    _set_cell_value(native, value)
    _set_cell_value(xml, value)
    assert native.Value2 == ("'"+value if isinstance(value, str) and value else value)
    assert xml.Value2 == value


def test_windows_export_slot_released_after_failure_and_serializes(monkeypatch):
    exporter = ExcelExporter(get_catalog())
    draft = build_draft(sample_payload(), get_catalog())
    import coc7_card.exporters.excel as module
    slot = threading.BoundedSemaphore(1)
    monkeypatch.setattr(module, 'WINDOWS_EXCEL_SLOT', slot)
    entered = threading.Event()
    release = threading.Event()
    active = 0
    def work(*args):
        nonlocal active
        active += 1
        assert active == 1
        entered.set()
        assert release.wait(5)
        active -= 1
        raise ExcelExportError('test failure')
    monkeypatch.setattr(exporter, '_export_windows', work)
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(exporter.export, draft)
        assert entered.wait(5)
        second = pool.submit(exporter.export, draft)
        release.set()
        for future in (first, second):
            with pytest.raises(ExcelExportError, match='test failure'):
                future.result(timeout=10)
    assert slot.acquire(blocking=False)
    slot.release()


native_excel = pytest.mark.skipif(sys.platform != 'win32' or os.environ.get('COC7_TEST_EXCEL') != '1',
                                  reason='Set COC7_TEST_EXCEL=1 on Windows with desktop Excel')


@native_excel
@pytest.mark.parametrize('year,code,custom', [(1920,'USD',False), (2026,'GBP',True)])
def test_native_excel_export_import_round_trip(year, code, custom, tmp_path):
    catalog = get_catalog()
    payload = sample_payload('=1+1' if custom else 'Windows测试', year, code)
    if custom:
        payload['occupation'] = {'is_custom':True, 'name':'民俗研究员', 'credit_min':0,
                                 'credit_max':99, 'point_formula':'EDU*2+INT*2'}
        payload['custom_skill_slots'] = [s.template_slot for s in catalog.skills[:8]]
    portrait = io.BytesIO()
    Image.new('RGB', (80,120), '#4c7084').save(portrait, format='PNG')
    before = hashlib.sha256(catalog.template_path.read_bytes()).hexdigest()
    with TestClient(app) as client:
        response = client.post('/api/export/excel', data={'draft_json':json.dumps(payload)},
                               files={'portrait':('portrait.png', portrait.getvalue(), 'image/png')} if custom else None)
    assert response.status_code == 200, response.text
    assert response.headers['X-Workbook-Sheet-Count'] == '14'
    (tmp_path/'native-export.xlsx').write_bytes(response.content)
    result = import_investigator(response.content, catalog)
    assert result['draft']['identity']['name'] == payload['identity']['name']
    assert result['draft']['identity']['player'] == '00123'
    assert result['draft']['background']['appearance'] == '=1+1'
    assert result['draft']['assets']['currency_code'] == code
    assert float(result['draft']['assets']['usd_cash']) == 123.45
    if custom:
        assert result['portrait'] is not None
        assert result['draft']['occupation_mode'] == 'custom'
    book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    try:
        assert tuple(book.sheetnames) == catalog.sheet_names
        assert book['更新说明']['P3'].value == catalog.template_revision_note
        rate = float(next(q['rate'] for q in catalog.currency_quotes if q['year']==year and q['code']==code))
        assert book['人物卡']['O62'].value == pytest.approx(round(123.45*rate, 2))
        # Excel omits the attribute when it uses the OOXML default, automatic.
        assert book.calculation.calcMode in (None, 'auto')
    finally:
        book.close()
    assert hashlib.sha256(catalog.template_path.read_bytes()).hexdigest() == before

    # Verify the saved file really recalculates after ordinary user input.
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    excel = None
    native = None
    fx = None
    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3
        native = excel.Workbooks.Open(str(tmp_path/'native-export.xlsx'), UpdateLinks=0, ReadOnly=True, CorruptLoad=0)
        assert excel.Calculation == -4105
        fx = native.Worksheets('货币汇率')
        assert fx.Range('A1:A2').EntireRow.Hidden
        assert fx.Range('A8:A27').EntireRow.Hidden
        assert not fx.Range('A35:A1156').EntireRow.Hidden
        assert not fx.Range('A1:T1').EntireColumn.Hidden
        assert fx.Range('U1:V1').EntireColumn.Hidden
        assert fx.Range('X1:AE1').EntireColumn.Hidden
        assert fx.Range('J3:K7').NumberFormat == ';;;'
        assert not fx.Range('A3:A7').EntireRow.Hidden
        assert fx.Range('C5').Validation.Type == 3
        assert fx.Range('C6').Validation.Type == 3
        quote = next(q for q in catalog.currency_quotes if q['year']==2026 and q['code']=='EUR')
        fx.Range('K1').Value2 = 2026
        native.Worksheets('人物卡').Range('S62').Value2 = quote['name']
        assert fx.Range('K5').Value2 == pytest.approx(float(quote['rate']))
        assert native.Worksheets('人物卡').Range('S62').Validation.Type == 3
    finally:
        if native is not None:
            native.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        fx = native = excel = None
        pythoncom.CoUninitialize()


@native_excel
def test_native_parallel_exports_keep_independent_characters():
    catalog = get_catalog()
    def export(name):
        result = ExcelExporter(catalog).export(build_draft(sample_payload(name), catalog))
        return import_investigator(result.data, catalog)['draft']['identity']['name']
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert list(pool.map(export, ['并发甲','并发乙'])) == ['并发甲','并发乙']
