"""Measure one operation in a fresh process; run before/after with the same Python."""
from __future__ import annotations

import argparse
import ctypes
import gc
import json
import sys
import time
import tracemalloc
from dataclasses import asdict
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from coc7_card.catalog import TemplateCatalog
from coc7_card.exporters.excel import ExcelExporter
from coc7_card.exporters.xlsx_template import TemplateWorkbook, worksheet_paths
from coc7_card.template_config import TEMPLATE_PATH


def peak_working_set_mb():
    if sys.platform != 'win32':
        import resource
        peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        return peak / (1024 ** 2 if sys.platform == 'darwin' else 1024)

    class Counters(ctypes.Structure):
        _fields_ = [('cb', ctypes.c_ulong), ('PageFaultCount', ctypes.c_ulong),
                    *[(name, ctypes.c_size_t) for name in (
                        'PeakWorkingSetSize', 'WorkingSetSize', 'QuotaPeakPagedPoolUsage',
                        'QuotaPagedPoolUsage', 'QuotaPeakNonPagedPoolUsage',
                        'QuotaNonPagedPoolUsage', 'PagefileUsage', 'PeakPagefileUsage')]]
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.GetCurrentProcess.restype = ctypes.c_void_p
    psapi = ctypes.WinDLL('psapi', use_last_error=True)
    psapi.GetProcessMemoryInfo.argtypes = [ctypes.c_void_p, ctypes.POINTER(Counters), ctypes.c_ulong]
    counters = Counters()
    counters.cb = ctypes.sizeof(counters)
    if not psapi.GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(counters), counters.cb):
        raise ctypes.WinError(ctypes.get_last_error())
    return counters.PeakWorkingSetSize / 1024 ** 2


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('operation', choices=['catalog', 'verify', 'template'])
    parser.add_argument('--workbook', type=Path, default=TEMPLATE_PATH)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--snapshot', type=Path)
    args = parser.parse_args()
    if args.operation == 'verify':
        with ZipFile(args.workbook) as archive:
            parts = {name: archive.read(name) for name in ('xl/workbook.xml', 'xl/_rels/workbook.xml.rels')}
        # Read the expected revision from the project template outside the measured
        # operation, just as the application's cached catalog does before export.
        template = load_workbook(TEMPLATE_PATH, read_only=True, data_only=False)
        try:
            revision_note = template['更新说明']['P3'].value
        finally:
            template.close()
        exporter = ExcelExporter(SimpleNamespace(
            sheet_names=tuple(worksheet_paths(parts)), template_revision_note=revision_note,
        ))
    gc.collect()
    tracemalloc.start()
    started = time.perf_counter()
    if args.operation == 'catalog':
        result = TemplateCatalog.load(args.workbook)
    elif args.operation == 'verify':
        result = exporter._verify_export(args.workbook)
    else:
        result = TemplateWorkbook(args.workbook)
        result.close()
    seconds = time.perf_counter() - started
    current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    report = {'operation': args.operation, 'seconds': round(seconds, 3),
              'python_peak_mb': round(peak / 1024 ** 2, 3),
              'python_retained_mb': round(current / 1024 ** 2, 3),
              'process_peak_mb': round(peak_working_set_mb(), 3),
              'python': sys.version.split()[0]}
    if args.snapshot and args.operation == 'catalog':
        args.snapshot.parent.mkdir(parents=True, exist_ok=True)
        args.snapshot.write_text(json.dumps(asdict(result), ensure_ascii=False, default=str, sort_keys=True), encoding='utf-8')
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(report))


if __name__ == '__main__':
    main()
