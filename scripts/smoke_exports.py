"""Check a running deployment and save representative export artifacts."""
import argparse
import base64
import io
import json
from pathlib import Path
import urllib.parse
import urllib.request
import warnings

from openpyxl import load_workbook
from pypdf import PdfReader


def run(base_url, output):
    output.mkdir(parents=True,exist_ok=True)
    def get(route):
        with urllib.request.urlopen(base_url+route,timeout=30) as response:
            return response.read()
    def export(kind,draft):
        body=urllib.parse.urlencode({'draft_json':json.dumps(draft,ensure_ascii=False)}).encode()
        request=urllib.request.Request(base_url+'/api/export/'+kind,data=body,
                                      headers={'Content-Type':'application/x-www-form-urlencoded'})
        with urllib.request.urlopen(request,timeout=180) as response:
            return response.read()
    health=json.loads(get('/api/health'));assert health['status']=='ok'
    assert 'COC7' in get('/').decode()
    catalog=json.loads(get('/api/bootstrap'))
    assert len(catalog['currencies'])==107
    assert sum(len(v) for v in catalog['currencies'].values())==1061
    assert sum(len(v) for v in catalog['currency_missing'].values())==60
    q=next(q for q in catalog['currencies']['2026'] if q['code']=='GBP')
    draft={'version':3,'identity':{'name':'年度汇率验证','age':32,'story_year':2026,'era':'现代'},
           'attributes':dict(STR=60,CON=65,SIZ=55,DEX=70,APP=50,INT=75,POW=60,EDU=70,Luck=55),
           'occupation_id':next(o['occupation_id'] for o in catalog['occupations'] if o['credit_min']==0),
           'experience':{'selection':'custom','name':'档案调查','skill_points':40},
           'skills':[{'template_slot':'F16','interest_points':2,'extra_final':12,'experience_points':17},
                     {'template_slot':'AB16','interest_points':3,'extra_final':7,'experience_points':11},
                     {'template_slot':'F20','extra_final':5,'experience_points':3}],
           'assets':{'currency_code':'GBP','usd_cash':'123.45','usd_assets':'0','usd_spending':''}}
    pdf_result=json.loads(export('pdf',draft))
    pdf=base64.b64decode(pdf_result['pdf_base64'])
    (output/'annual-rate-example.pdf').write_bytes(pdf)
    reader=PdfReader(io.BytesIO(pdf))
    text=''.join(p.extract_text() for p in reader.pages)
    assert len(reader.pages)==2 and '年度汇率验证' in text
    assert '91.72' in text and '2026-08-31' in text and 'FRED' in text
    for i,preview in enumerate(pdf_result['previews'],1):
        (output/f'pdf-page-{i}.png').write_bytes(base64.b64decode(preview))
    xlsx=export('excel',draft)
    (output/'annual-rate-example.xlsx').write_bytes(xlsx)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore',UserWarning)
        book=load_workbook(io.BytesIO(xlsx),data_only=True)
    try:
        assert len(book.sheetnames)==14
        assert book['人物卡']['E3'].value=='年度汇率验证'
        assert book['人物卡']['O62'].value==91.72
        assert book['人物卡']['S62'].value==q['name']
        assert abs(book['货币汇率']['K5'].value-float(q['rate']))<1e-14
        assert book['货币汇率']['K7'].value==q['source']
        assert book['货币汇率']['K6'].value==q['date']
        assert book['人物卡']['L15'].value.replace('\n','')=='成长点数'
        assert book['人物卡']['M15'].value.replace('\n','')=='经历包点'
        for cell,expected in {'L16':12,'M16':17,'R16':36,'T16':18,'V16':7,
                              'AH16':7,'AI16':11,'AN16':26,'L20':5,'M20':3,'R20':13}.items():
            assert book['人物卡'][cell].value==expected,(cell,book['人物卡'][cell].value)
        assert book['附表']['AC27'].value==9
        assert '剩余经历包点=9' in book['人物卡']['J50'].value
        errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e']
        assert not errors,errors
    finally: book.close()
    result={'health':health,'years':107,'available_quotes':1061,'missing_quotes':60,
            'pdf_pages':len(reader.pages),'excel_sheets':14,'cash_gbp':'91.72',
            'growth_points':24,'experience_points':31,'experience_remaining':9,
            'template_sha256':catalog['meta']['source_sha256'],
            'source_period':q['date'],'source':q['source'],'data_sha256':catalog['currency_metadata']['data_sha256']}
    (output/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url',default='http://127.0.0.1:8765')
    parser.add_argument('--output',type=Path,default=Path('test-output/deployment-smoke'))
    args=parser.parse_args()
    run(args.url.rstrip('/'),args.output)
