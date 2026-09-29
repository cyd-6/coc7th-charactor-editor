"""Import the supplied workbook without changing it or inventing missing rates.

Run with the bundled Python runtime (openpyxl is only used for reading).
The annual matrix is authoritative; raw records supply provenance and checks.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import warnings
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_NAME = "historical_exchange_rates_1920_2026.xlsx"
SLOTS = ["USD", "GBP", "CAD", "FR", "DE", "IT", "JPY", "CHF", "INR", "RUB", "CNY", "EUR"]


def currency_code(slot, year):
    if slot == "FR":
        return "FRF_OLD" if year < 1960 else "FRF"
    if slot == "DE":
        if year <= 1922:
            return "DRM"
        if year in (1923, 1924, 1948):
            return f"DE_TRANSITION_{year}"
        return "RM" if year < 1948 else "DEM"
    if slot == "IT":
        return "ITL"
    if slot == "RUB":
        for end, code in [(1921,"RUB_EARLY"),(1922,"RUB_1922"),(1923,"RUB_1923"),
                          (1946,"SUR_1924"),(1947,"SUR_TRANSITION_1947"),(1960,"SUR_1947"),
                          (1991,"SUR_1961"),(1997,"RUR")]:
            if year <= end:
                return code
    if slot == "CNY":
        for end, code in [(1921,"TAEL"),(1932,"CN_SILVER_OLD"),(1934,"CN_SILVER"),
                          (1935,"CN_TRANSITION_1935"),(1947,"FABI"),(1948,"CN_TRANSITION_1948"),
                          (1954,"CNY_OLD")]:
            if year <= end:
                return code
    return slot


def text(value):
    return "" if value is None else str(value)


def annual_numeric_text(path):
    """Keep the decimal representation saved by the source workbook."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(path) as archive:
        sheets = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        sheet = next(s for s in sheets.find("m:sheets", ns) if s.get("name") == "年度汇率")
        target = rels[sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")]
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        root = ET.fromstring(archive.read(part))
        return {c.get("r"): c.find("m:v", ns).text for c in root.findall(".//m:c", ns)
                if c.get("t") not in ("str", "s", "e", "inlineStr") and c.find("m:v", ns) is not None
                and c.find("m:v", ns).text}


def import_workbook(path):
    decimals = annual_numeric_text(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        book = load_workbook(path, read_only=True, data_only=True)
    try:
        matrices = {name: list(book[name].iter_rows(min_row=6, max_row=112, values_only=True))
                    for name in ("年度汇率", "实际币名", "口径与来源")}
        sources = {}
        for row in book["来源目录"].iter_rows(min_row=6, values_only=True):
            if len(row) >= 5 and isinstance(row[0], str) and row[0] and row[2] and str(row[2]).startswith("http"):
                if row[0] in sources:
                    raise ValueError(f"重复的来源编号：{row[0]}")
                sources[row[0]] = row
        raw = {}
        for n, row in enumerate(book["原始记录"].iter_rows(min_row=6, values_only=True), 6):
            if len(row) < 13 or not isinstance(row[0], (int, float)):
                continue
            key = (int(row[0]), row[1])
            if key in raw:
                raise ValueError(f"重复的原始报价：{key}")
            raw[key] = (row, n)
        missing = {(int(r[0]), r[1]): text(r[3]) for r in book["缺失清单"].iter_rows(min_row=6, values_only=True)
                   if len(r) >= 4 and isinstance(r[0], (int, float))}
        quotes = []
        for index, annual in enumerate(matrices["年度汇率"]):
            year, rownum = int(annual[0]), index + 6
            assert year == 1920 + index
            for col, slot in enumerate(SLOTS, 2):
                if (year < 1999 and slot == "EUR") or (year >= 1999 and slot in ("FR", "DE", "IT")):
                    continue
                cell = f"{chr(64+col)}{rownum}"
                rate = decimals.get(cell)
                name = text(matrices["实际币名"][index][col-1])
                source_cells = [f"年度汇率!{cell}"]
                if slot == "DE" or (slot == "EUR" and year >= 1999):
                    duplicate_cols = [13] if slot == "DE" else [5, 6, 7]
                    for duplicate in duplicate_cols:
                        dup_cell = f"{chr(64+duplicate)}{rownum}"
                        assert decimals.get(dup_cell) == rate, (year, slot, dup_cell)
                        source_cells.append(f"年度汇率!{dup_cell}")
                quote = dict(year=year, code=currency_code(slot, year), slot=slot, name=name,
                             available=rate is not None, rate=rate, date=str(year), basis="缺失／未核实",
                             source="用户提供的历史汇率表", source_id="", source_url="", source_grade="",
                             notes="", conversion="", raw_quote="", raw_direction="", scale="",
                             source_cells=source_cells, raw_cell="", missing_reason="", warning="")
                if slot == "USD":
                    assert Decimal(rate) == 1
                    quote.update(basis="基准恒等", notes="1 美元等于 1 美元，无外部报价。", conversion="美元基准，汇率恒为 1。")
                elif rate is None:
                    assert (year, slot) not in raw
                    quote["missing_reason"] = missing.get((year, slot), "来源表未提供可核实报价；未插补。")
                else:
                    record, record_row = raw[(year, slot)]
                    assert Decimal(rate) > 0
                    assert abs(Decimal(str(record[6]))-Decimal(rate)) <= abs(Decimal(rate))*Decimal("1e-14")
                    source = sources[record[9]]
                    raw_value, scale = Decimal(str(record[3])), Decimal(str(record[5]))
                    converted = raw_value * scale if record[4] == "LCU/USD" else scale / raw_value
                    assert abs(converted-Decimal(rate)) <= abs(Decimal(rate))*Decimal("1e-14")
                    method = "原报价已按 1 美元兑换本币表示。" if record[4] == "LCU/USD" else "原报价为每单位本币兑美元，取原报价均值的倒数；不等于反向日度汇率的算术平均。"
                    if scale != 1:
                        method += f"乘以 {scale}，还原当年币制的计价单位。"
                    quote.update(date=text(record[8]), basis=text(record[7]), source=text(source[1]),
                                 source_id=text(record[9]), source_url=text(record[10] or source[2]),
                                 source_grade=text(record[12]), notes="\n".join(dict.fromkeys(filter(None,[text(record[11]),text(source[3])]))),
                                 conversion=method, raw_quote=str(raw_value), raw_direction=text(record[4]), scale=str(scale),
                                 raw_cell=f"原始记录!A{record_row}:M{record_row}")
                    cautions = []
                    if "官价" in quote["basis"]:
                        cautions.append("官价参考，并非市场年均汇率")
                    elif "部分" in quote["basis"] or "年内" in quote["basis"]:
                        cautions.append("仅覆盖部分期间")
                    if year == 2026:
                        cautions.append("2026 年仅为 1—8 月均值，并非全年均值")
                    if "过渡" in name or "并行" in name or "／" in name:
                        cautions.append("币制过渡期，请留意实际币名与观测日期")
                    quote["warning"] = "；".join(cautions)
                quotes.append(quote)
        assert len(quotes) == 1121 and sum(q["available"] for q in quotes) == 1061
        assert len({(q["year"], q["code"]) for q in quotes}) == len(quotes)
        return dict(schema_version=1, source_file=path.name, source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                    year_start=1920, year_end=2026, base_currency="USD", rate_unit="每 1 美元兑换的当年本币数量",
                    note="数据来自用户提供的历史汇率表，缺失值未插补，历史报价的统计口径可能不同。",
                    quotes=quotes)
    finally:
        book.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "assets/rates/sources" / SOURCE_NAME)
    parser.add_argument("--output", type=Path, default=ROOT / "assets/rates/annual.json")
    args = parser.parse_args()
    dataset = import_workbook(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"导入 107 年，{len(dataset['quotes'])} 条记录，1061 条有效报价，60 条缺失。")
