"""Preserve native OOXML while inserting the Artifact Tool authored ranges."""
from __future__ import annotations
import argparse
import copy
import json
import sys
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from lxml import etree as ET
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from coc7_card.exporters.xlsx_template import TemplateWorkbook, tag, worksheet_paths


def insert_worksheet_control(root, control):
    """Insert controls in CT_Worksheet order, after merges and phoneticPr.

    XML parsers accept out-of-order children, but Excel repairs the sheet.
    Keep every unrelated child in its original position.
    """
    predecessors = {
        tag(name) for name in (
            'sheetPr', 'dimension', 'sheetViews', 'sheetFormatPr', 'cols',
            'sheetData', 'sheetCalcPr', 'sheetProtection', 'protectedRanges',
            'scenarios', 'autoFilter', 'sortState', 'dataConsolidate',
            'customSheetViews', 'mergeCells', 'phoneticPr',
            'conditionalFormatting',
        )
    }
    if control.tag not in {tag('conditionalFormatting'), tag('dataValidations')}:
        raise ValueError('Expected a worksheet validation or conditional format')
    if control.tag == tag('dataValidations'):
        predecessors.add(tag('dataValidations'))
    position = max((i + 1 for i, child in enumerate(root)
                    if child.tag in predecessors), default=0)
    root.insert(position, control)


def discard_formula_cache(cell):
    """Native calculation owns caches, including unsupported donor functions."""
    if cell.find(tag('f')) is not None:
        for child in list(cell):
            if child.tag in {tag('v'),tag('is')}: cell.remove(child)
        cell.attrib.pop('t',None)


def merge(template, donor_path, manifest_path, output):
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    calculator = manifest.get('upgrade') == 'cross-v1'
    compact = manifest.get('upgrade') == 'compact-v1'
    with ZipFile(donor_path) as z:
        donor = {name:z.read(name) for name in z.namelist()}
    paths = worksheet_paths(donor)
    strings = []
    if 'xl/sharedStrings.xml' in donor:
        strings = list(ET.fromstring(donor['xl/sharedStrings.xml']))
    book = TemplateWorkbook(template)
    try:
        original = book.tree('xl/styles.xml')
        incoming = ET.fromstring(donor['xl/styles.xml'])
        maps = {}
        for element in ('fonts','fills','borders','dxfs'):
            source = incoming.find(tag(element))
            if source is None:
                maps[element] = {}; continue
            target = original.find(tag(element))
            if target is None:
                target = ET.Element(tag(element),count='0')
                before = original.find(tag('tableStyles'))
                original.insert(list(original).index(before) if before is not None else len(original),target)
            maps[element] = {}
            for i,node in enumerate(source):
                maps[element][i] = len(target)
                target.append(copy.deepcopy(node))
            target.set('count',str(len(target)))
        number_map = {}
        numbers = original.find(tag('numFmts'))
        if numbers is None:
            numbers = ET.Element(tag('numFmts'),count='0');original.insert(0,numbers)
        source_numbers = incoming.find(tag('numFmts'))
        if source_numbers is not None:
            for node in source_numbers:
                match = next((n for n in numbers if n.get('formatCode')==node.get('formatCode')),None)
                if match is None:
                    new_id = max([163, *(int(n.get('numFmtId')) for n in numbers)])+1
                    match = copy.deepcopy(node);match.set('numFmtId',str(new_id));numbers.append(match)
                number_map[int(node.get('numFmtId'))] = int(match.get('numFmtId'))
        numbers.set('count',str(len(numbers)))
        formats = original.find(tag('cellXfs'))
        style_map = {}
        for i,node in enumerate(incoming.find(tag('cellXfs'))):
            cloned = copy.deepcopy(node)
            for attr, group in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:
                if cloned.get(attr) is not None:
                    cloned.set(attr,str(maps[group][int(cloned.get(attr))]))
            num = int(cloned.get('numFmtId','0'))
            cloned.set('numFmtId',str(number_map.get(num,num)))
            cloned.set('xfId','0')
            style_map[i] = len(formats);formats.append(cloned)
        formats.set('count',str(len(formats)))

        if calculator:
            prepare_calculator_layout(book, manifest)
        for name in (('货币汇率',) if compact else ('货币汇率','人物卡') if calculator else ('货币汇率','人物卡','附表')):
            source = ET.fromstring(donor[paths[name]])
            sheet = book.Worksheets(name)
            for cell in source.iter(tag('c')):
                address = cell.get('r')
                if compact and address not in manifest['cells']:
                    continue
                if name == '人物卡' and address not in manifest['cardCells']:
                    continue
                if name == '附表' and address not in manifest['appendixCells']:
                    continue
                row_index, column_index = coordinate_to_tuple(address)
                if calculator and name == '货币汇率' and not (
                    row_index <= 31 and column_index <= 6
                    or address in {'K2','K24','U3','V3','U4','V4'}
                ):
                    continue
                target = sheet.cell(column_index, row_index)
                style = target.get('s','0')
                for child in list(target): target.remove(child)
                target.attrib.clear();target.set('r',address)
                for key,value in cell.attrib.items(): target.set(key,value)
                # Keep the card's existing appearance; style only new FX areas.
                use_donor_style = not compact and name == '货币汇率' and (
                    column_index <= 6 if calculator else address != 'A1')
                target.set('s',str(style_map[int(cell.get('s','0'))]) if use_donor_style else style)
                if compact and address == 'C4':
                    overlay = copy.deepcopy(formats[int(style)])
                    overlay.set('numFmtId','0');overlay.set('applyNumberFormat','1')
                    target.set('s',str(len(formats)));formats.append(overlay)
                    formats.set('count',str(len(formats)))
                if calculator and (name == '人物卡' and address in {'B60','S62'} or name == '货币汇率' and address == 'K2'):
                    # Change only the input/formula cue, retaining the card's
                    # native border, font, number format and merged alignment.
                    overlay = copy.deepcopy(formats[int(style)])
                    if address != 'B60':
                        overlay.set('fillId',formats[style_map[int(cell.get('s','0'))]].get('fillId'))
                    if name == '人物卡':
                        alignment = overlay.find(tag('alignment'))
                        if alignment is None: alignment = ET.SubElement(overlay,tag('alignment'))
                        alignment.set('wrapText','0'); alignment.set('shrinkToFit','1')
                        overlay.set('applyAlignment','1')
                    target.set('s',str(len(formats)))
                    formats.append(overlay); formats.set('count',str(len(formats)))
                if cell.get('t')=='s':
                    si = strings[int(cell.find(tag('v')).text)]
                    inline = copy.deepcopy(si);inline.tag=tag('is')
                    target.set('t','inlineStr');target.append(inline)
                else:
                    for child in cell: target.append(copy.deepcopy(child))
                # HYPERLINK may emit a non-OOXML error string in the donor.
                # Recalculate in native Excel/LibreOffice before replacement.
                discard_formula_cache(target)
            if compact:
                prepare_compact_layout(sheet, source, maps)
                continue
            if calculator:
                merge_calculator_controls(sheet, source, maps, name)
                if name == '货币汇率':
                    for row in source.find(tag('sheetData')):
                        n = int(row.get('r'))
                        if n <= 27 and row.get('ht'):
                            height = float(row.get('ht'))
                            # The right-hand asset source notes still share
                            # these rows; preserve their existing readable height.
                            if 14 <= n <= 24:
                                height = max(height,float(sheet.rows[n].get('ht','15')))
                            sheet.rows[n].set('ht',str(height))
                            sheet.rows[n].set('customHeight','1')
                            sheet.rows[n].attrib.pop('hidden',None)
                continue
            if name != '货币汇率': continue
            for row in source.find(tag('sheetData')):
                n = int(row.get('r'))
                if n not in sheet.rows: continue
                target = sheet.rows[n]
                if row.get('ht'):
                    target.set('ht',str(max(float(target.get('ht','15')),float(row.get('ht')))))
                    target.set('customHeight','1')
                if n==2: target.attrib.pop('hidden',None)
            # The currency selector is on row 2, hidden in the legacy sheet.
            sheet.rows[2].attrib.pop('hidden',None)
            view=sheet.root.find(tag('sheetViews'))
            if view is not None and len(view):
                view[0].set('topLeftCell','J1')
                selection=view[0].find(tag('selection'))
                if selection is not None:
                    selection.set('activeCell','K1');selection.set('sqref','K1')
            for col in source.find(tag('cols')):
                start,end = int(col.get('min')),int(col.get('max'))
                for n in range(max(11,start),end+1):
                    sheet.Columns(get_column_letter(n)).ColumnWidth = float(col.get('width','16'))
            for kind in ('dataValidations','conditionalFormatting'):
                # Re-running the import replaces our controls, retaining any
                # other validations and conditional rules in the template.
                for existing in list(sheet.root.findall(tag(kind))):
                    if kind=='conditionalFormatting' and existing.get('sqref') in ('K8','K8:K8'):
                        sheet.root.remove(existing)
                    elif kind=='dataValidations':
                        for rule in list(existing):
                            if rule.get('sqref') in ('K1','K2'): existing.remove(rule)
                        if not len(existing): sheet.root.remove(existing)
                        else: existing.set('count',str(len(existing)))
                for item in source.findall(tag(kind)):
                    cloned=copy.deepcopy(item)
                    for rule in cloned.iter(tag('cfRule')):
                        if rule.get('dxfId') is not None: rule.set('dxfId',str(maps['dxfs'][int(rule.get('dxfId'))]))
                    for validation in cloned.iter(tag('dataValidation')):
                        validation.set('showErrorMessage','1');validation.set('errorStyle','stop')
                        validation.set('errorTitle','请选择有效的年份和币种')
                        validation.set('error','年份为1920—2026；币种须从当前年份下拉列表选择。')
                        validation.set('showDropDown','0')
                    if kind=='dataValidations':
                        existing=sheet.root.find(tag(kind))
                        if existing is not None:
                            existing.extend(cloned)
                            existing.set('count',str(len(existing)))
                            continue
                    insert_worksheet_control(sheet.root,cloned)
            dimension=sheet.root.find(tag('dimension'))
            if dimension is not None: dimension.set('ref',f'A1:V{manifest["last"]}')
        workbook=book.tree('xl/workbook.xml')
        names=workbook.find(tag('definedNames'))
        if names is None:
            names=ET.Element(tag('definedNames'))
            workbook.insert(list(workbook).index(workbook.find(tag('sheets')))+1,names)
        for record in manifest['names']:
            for old in list(names):
                if old.get('name')==record['name']: names.remove(old)
            ET.SubElement(names,tag('definedName'),name=record['name']).text=record['formula']
        output.parent.mkdir(parents=True,exist_ok=True)
        book.save(output)
    finally:
        book.close()


def prepare_calculator_layout(book, manifest):
    """Move only the former reference block; never regenerate user data."""
    sheet = book.Worksheets('货币汇率')
    names = book.tree('xl/workbook.xml').find(tag('definedNames'))
    upgraded = names is not None and any(n.get('name') == 'COC7_FX_CALCULATOR_SCHEMA' for n in names)
    if not upgraded:
        # The reference contains only values, no formulas or drawing anchors.
        # Fail rather than overwrite a user's new data in the destination.
        moving = [(address,copy.deepcopy(cell)) for address,cell in sheet.cells.items()
                  if (lambda rc: rc[0] <= 31 and rc[1] <= 8)(coordinate_to_tuple(address))]
        for address,cell in moving:
            row,col = coordinate_to_tuple(address)
            destination = f'{get_column_letter(col+23)}{row+35}'
            existing = sheet.cells.get(destination)
            if existing is not None and len(existing):
                raise ValueError(f'历史参考目标位置已有内容：{destination}')
            if cell.find(tag('f')) is not None:
                raise ValueError(f'历史参考含公式，须先核对引用：{address}')
        for address,cell in moving:
            row,col = coordinate_to_tuple(address)
            target = sheet.cell(col+23,row+35)
            target.attrib.clear(); target.attrib.update(cell.attrib)
            target.set('r',f'{get_column_letter(col+23)}{row+35}')
            for child in list(target): target.remove(child)
            for child in cell: target.append(copy.deepcopy(child))
            original = sheet.cells[address]
            original.getparent().remove(original)
            del sheet.cells[address]
        # Preserve the legacy reference's column widths and enough row height.
        for col in range(1,9):
            width = next((float(node.get('width')) for node in sheet.root.find(tag('cols'))
                          if int(node.get('min')) <= col <= int(node.get('max'))),13)
            sheet.Columns(get_column_letter(col+23)).ColumnWidth = width
        for row in range(1,32):
            old = sheet.rows.get(row); new = sheet.rows.get(row+35)
            if old is not None and new is not None and old.get('ht'):
                new.set('ht',str(max(float(old.get('ht')),float(new.get('ht','15')))))
                new.set('customHeight','1')
    merges = sheet.root.find(tag('mergeCells'))
    if merges is None:
        raise ValueError('模板缺少原始合并区域')
    owned = set(manifest['merges']) | {'A1:H1'}
    for node in list(merges):
        if node.get('ref') in owned: merges.remove(node)
    existing = {node.get('ref') for node in merges}
    for area in [*manifest['merges'],'X36:AE36']:
        if area not in existing: ET.SubElement(merges,tag('mergeCell'),ref=area)
    merges.set('count',str(len(merges)))
    view = sheet.root.find(tag('sheetViews'))[0]
    view.set('topLeftCell','A1')
    selection = view.find(tag('selection'))
    if selection is not None:
        selection.set('activeCell','C4'); selection.set('sqref','C4')
    sheet.root.find(tag('dimension')).set('ref',f'A1:AE{manifest["last"]}')
    sheet.PageSetup.PrintArea = '$A$1:$F$27'


def prepare_compact_layout(sheet, source, maps):
    """Expose only the five existing controls, retaining every backing cell."""
    for n,row in sheet.rows.items():
        if n < 3 or n > 7: row.set('hidden','1')
        else:
            row.attrib.pop('hidden',None)
            row.set('ht','42');row.set('customHeight','1')
    for col in sheet.root.find(tag('cols')):
        if int(col.get('min')) >= 7: col.set('hidden','1')
    # G:AE holds all right-hand sources and the existing asset controls.
    # Blank columns after it need no XML styles or million-cell allocations.
    view = sheet.root.find(tag('sheetViews'))[0]
    view.set('topLeftCell','A3');view.set('showGridLines','0')
    selection = view.find(tag('selection'))
    if selection is not None:
        selection.set('activeCell','C4');selection.set('sqref','C4')
    for existing in list(sheet.root.findall(tag('conditionalFormatting'))):
        if existing.get('sqref') == 'C7:F7': sheet.root.remove(existing)
    for control in source.findall(tag('conditionalFormatting')):
        cloned = copy.deepcopy(control)
        for rule in cloned.iter(tag('cfRule')):
            if rule.get('dxfId') is not None:
                rule.set('dxfId',str(maps['dxfs'][int(rule.get('dxfId'))]))
        insert_worksheet_control(sheet.root,cloned)
    sheet.PageSetup.PrintArea = '$A$3:$F$7'


def merge_calculator_controls(sheet, source, maps, name):
    owned_validations = {'K2','C3','C5','C6'} if name == '货币汇率' else {'S62'}
    for container in list(sheet.root.findall(tag('dataValidations'))):
        for rule in list(container):
            if rule.get('sqref') in owned_validations: container.remove(rule)
        if not len(container): sheet.root.remove(container)
        else: container.set('count',str(len(container)))
    for old in list(sheet.root.findall(tag('conditionalFormatting'))):
        if old.get('sqref') == 'C9:F9': sheet.root.remove(old)
    for kind in ('conditionalFormatting','dataValidations'):
        for control in source.findall(tag(kind)):
            cloned = copy.deepcopy(control)
            for rule in cloned.iter(tag('cfRule')):
                if rule.get('dxfId') is not None:
                    rule.set('dxfId',str(maps['dxfs'][int(rule.get('dxfId'))]))
            for rule in cloned.iter(tag('dataValidation')):
                rule.set('showDropDown','0'); rule.set('showErrorMessage','1')
                rule.set('errorStyle','stop'); rule.set('allowBlank','0')
                rule.set('showInputMessage','1')
                rule.set('promptTitle','币种与年份')
                rule.set('prompt','年份改变后，请重新选择当年有报价的币种。主表资产年份在货币汇率 K1 修改。')
                rule.set('errorTitle','输入无效')
                rule.set('error','年份须为1920—2026的整数；币种须从当年下拉列表选择。')
            existing = sheet.root.find(tag(kind)) if kind == 'dataValidations' else None
            if existing is not None:
                existing.extend(cloned); existing.set('count',str(len(existing)))
            else:
                insert_worksheet_control(sheet.root,cloned)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--donor',type=Path,default=ROOT/'test-output/fx-upgrade/fx-donor.xlsx')
    parser.add_argument('--manifest',type=Path,default=ROOT/'test-output/fx-upgrade/fx-manifest.json')
    args=parser.parse_args()
    merge(args.template,args.donor,args.manifest,args.output)
    print(args.output)
